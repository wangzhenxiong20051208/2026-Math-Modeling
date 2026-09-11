# -*- coding: utf-8 -*-
r"""
2026 年高教社杯全国大学生数学建模竞赛  C 题
微网与外部电网电力调控策略 —— 问题二

考虑预测误差与紧急购电成本的日前计划 + 实时执行模型
--------------------------------------------------------------------------
实现依据：《问题二_解题思路与实现框架》
  01_题目/C题/问题二_解题思路与实现框架.md

口径（与框架文档一致，并与问题一保持同一套符号）：
  * 内部按当天 0:00-24:00 划分 144 个时段，Δt = 1/6 h。
  * 附件 2 的列标签同样采用区间**结束时刻**口径（首列 0:10、末列 0:00+1），
    故第 j 列对应区间 [(j-1)*10min, j*10min)。
  * 决策变量一律用**电量**（kWh）：g_t 计划购电量、c_t 充电量、d_t 放电量、
    w_t 弃用电量、E_t 时段末储电量（145 个状态点）、z_t ∈ {0,1} 模式变量。
  * 充放电为母线侧口径：充入 1 kWh 电池增加 η_c = 0.9 kWh；
    放出 1 kWh 电池减少 1/η_d = 1/0.9 kWh；往返效率 0.81。
  * 费用 = Σ p_t g_t（计划费，不论是否用完全额收取）
         + Σ 5 p_t r_t（紧急费，量在计划之外）。

三层结构：
  1) 预测层 —— 仅用当天之前的历史负载/光伏做加权预测，得到净负荷预测；
  2) 计划层 —— 用历史预测误差的经验 α 分位数修正净负荷，求解日 MILP，
                冻结当天 144 段计划购电量 g；
  3) 执行层 —— 逐段观测实际负载/光伏，按因果反馈规则决定充放电、
                计算紧急购电与弃电，并把实际日末储电量传给次日。

时间纪律：任何时段的 g_{n,t} 都不得在看到当天实际值后被修改；
预测、分位数、参数一律只使用当前日期之前的数据。
--------------------------------------------------------------------------
"""

from __future__ import annotations

import argparse
import json
import sys
import time
from dataclasses import dataclass, field, asdict
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import csr_matrix

sys.path.insert(0, str(Path(__file__).resolve().parent))
from p1_microgrid import (  # noqa: E402
    N, TAU, E_MIN, E_MAX, E_INIT, P_MAX, M_ENERGY, ETA,
    interval_label, block_label, fmt_time, BLOCKS,
)

# ---------------------------------------------------------------- 路径
ROOT = Path(__file__).resolve().parents[1]
ATTACH2 = ROOT / "01_题目" / "C题" / "附件" / "附件2.xlsx"
TEMPLATE2 = ROOT / "01_题目" / "C题" / "附件" / "附件5" / "result2.xlsx"
OUT_DIR = ROOT / "06_支撑材料"
OUT_XLSX = OUT_DIR / "result2.xlsx"
OUT_JSON = OUT_DIR / "p2_results.json"
OUT_DETAIL = OUT_DIR / "p2_detail.csv"
OUT_DAILY = OUT_DIR / "p2_daily.csv"
OUT_CAL = OUT_DIR / "p2_calibration.csv"

N_DAY = 365                  # 全年天数
DAY0 = np.datetime64("2025-01-01")

# 目标函数中的建模选择（框架文档 4.1 节）
E_TAR = float(E_INIT)        # 希望为次日保留的参考储电量 kWh
LAMBDA_DEFAULT = 0.60        # 日末储备不足的规划代价 元/kWh

# 正式输出区间：2025-02-01 至 2025-12-31，共 334 天
REPORT_START = 31
REPORT_END = 365

# 指定展示日期（题目表 3）
SPECIAL_DATES = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]

WANT_LABELS = ["10:00-10:10", "12:00-12:10", "14:00-14:10",
               "16:00-16:10", "18:00-18:10", "20:00-20:10"]

# 紧急购电事件合并阈值（仅过滤浮点噪声，kWh）
EMG_EPS = 1e-6


# ================================================================ 数据层
def load_attach2() -> tuple[np.ndarray, np.ndarray, list[str]]:
    """读取附件 2 的全年小区负载与光伏实际功率。

    返回 (LOAD, PV, dates)，形状均为 (365, 144)，单位 kW；dates 为 'YYYY-MM-DD'。
    同时校验数据完整性——框架文档 2.1 节要求保留这些基础检查。
    """
    df = pd.read_excel(ATTACH2, sheet_name="小区负载")
    dfp = pd.read_excel(ATTACH2, sheet_name="光伏发电实际功率")

    if df.shape != dfp.shape:
        raise ValueError(f"附件 2 两张表形状不一致：{df.shape} vs {dfp.shape}")
    # pandas 已消费表头行，故数据体应为 365 天 × (日期列 + 144 个时段)
    if df.shape != (N_DAY, N + 1):
        raise ValueError(f"附件 2 形状应为 {(N_DAY, N + 1)}，实际 {df.shape}")

    dates = [pd.Timestamp(v).strftime("%Y-%m-%d") for v in df.iloc[:, 0]]
    if len(set(dates)) != N_DAY:
        raise ValueError("附件 2 日期不唯一")
    if not np.all(np.diff(np.array([int(d.replace('-', '')) for d in dates])) > 0):
        raise ValueError("附件 2 日期非严格递增")
    if dates[0] != "2025-01-01" or dates[-1] != "2025-12-31":
        raise ValueError(f"附件 2 日期区间应为 2025-01-01~2025-12-31，实际 "
                         f"{dates[0]}~{dates[-1]}")

    LOAD = df.iloc[:, 1:].to_numpy(float)
    PV = dfp.iloc[:, 1:].to_numpy(float)
    if LOAD.shape != (N_DAY, N) or PV.shape != (N_DAY, N):
        raise ValueError(f"附件 2 数据体形状应为 {(N_DAY, N)}")

    if np.isnan(LOAD).any() or np.isnan(PV).any():
        raise ValueError("附件 2 存在空值")
    if LOAD.min() < 0 or PV.min() < 0:
        raise ValueError("附件 2 存在负功率值")

    # 末列必须是 0:00+1，否则区间口径不成立
    last = df.columns[-1]
    if str(last).strip() != "0:00+1":
        raise ValueError(f"附件 2 末列表头应为 0:00+1，实际 {last!r}")

    return LOAD, PV, dates


def weekday_of(n: int) -> int:
    """第 n 天（0 起，2025-01-01 为周三）的星期序号，0 = 周一。"""
    # 2025-01-01 是星期三，weekday() = 2
    return (2 + n) % 7


# ================================================================ 预测层
@dataclass(frozen=True)
class ForecastParams:
    """近期加权历史预测的超参数（框架文档 3.2 节）。

    权重形式：a_i ∝ decay^((n-1-i)/7)，即按"周"衰减；weekday_only 为真时
    只保留与待预测日同一星期的历史日（同星期日的结构在本题中极强）。
    只使用 i < n 的历史，天然满足因果性。
    """

    window: int = 15          # 回看窗口天数
    decay: float = 0.30       # 每 7 天的权重衰减因子
    weekday_only: bool = False  # 是否只用同一星期的历史
    name: str = "负荷"

    def label(self) -> str:
        wp = "仅同星期" if self.weekday_only else "全部历史"
        return f"{self.name}：近{self.window}天/{wp}/周衰减{self.decay:g}"


# 在正式区间开始前选定并冻结的预测超参数，回放中不再调整。
# 结构依据：同族内各工作日的日内形状与族平均形状高度一致——把每个「星期几」的
# 平均形状与其所属族的平均形状相比，逐 (星期几, 时段) 相对偏差的中位数约 0.25%
# （族A 0.250%、族B 0.254%；最大约 22%，出现在个别时段）。
# 注意别与「族间」差异混淆：族间同一口径的中位数是 8.9%、最大 38.5%，那是论文
# §6.3 引用的数字，量级完全不同。
# 日总量按星期分成两族（周一至周四、周日约为周五、周六的 1.58 倍），故负载须按
# 星期分组预测；光伏无星期结构但有强季节性，故用最近数日。
# 注：本注释下方论文 §6.3 中的对照数字在正式区间上统计，属对该结构选择的事后
# 评价；仅用 1 月数据判定同样成立（净负荷日绝对误差 20173.9 → 5903.7 kWh/日）。
LOAD_FORECAST = ForecastParams(window=15, decay=0.30, weekday_only=True, name="负荷")
PV_FORECAST = ForecastParams(window=5, decay=0.50, weekday_only=False, name="光伏")


def forecast_weights(n: int, fp: ForecastParams) -> tuple[np.ndarray, np.ndarray]:
    """返回第 n 天预测所用的 (历史日下标, 归一化权重)。

    同星期历史为空时（前 7 天）退化为窗口内的全部历史日，保证总有可用预测。
    """
    lo = max(0, n - fp.window)
    idx = np.arange(lo, n)
    if len(idx) == 0:
        return idx, np.zeros(0)
    w = fp.decay ** ((n - 1 - idx) / 7.0)
    if fp.weekday_only:
        same = np.array([weekday_of(i) == weekday_of(n) for i in idx])
        if same.any():
            w = w * same
    s = w.sum()
    if s <= 0:
        return idx, np.full(len(idx), 1.0 / len(idx))
    return idx, w / s


class Forecaster:
    """因果预测器：第 n 天的预测只依赖第 n 天之前的历史。"""

    def __init__(self, LOAD: np.ndarray, PV: np.ndarray,
                 fp_load: ForecastParams = LOAD_FORECAST,
                 fp_pv: ForecastParams = PV_FORECAST):
        self.LOAD, self.PV = LOAD, PV
        self.fp_load, self.fp_pv = fp_load, fp_pv
        self._wl = [forecast_weights(n, fp_load) for n in range(N_DAY)]
        self._wv = [forecast_weights(n, fp_pv) for n in range(N_DAY)]

    def raw(self, n: int) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
        """返回 (负载预测 kW, 光伏预测 kW, 负载历史日下标, 光伏历史日下标)。"""
        il, wl = self._wl[n]
        iv, wv = self._wv[n]
        if len(il) == 0 or len(iv) == 0:
            raise ValueError(f"第 {n} 天没有可用历史，应走冷启动分支")
        return wl @ self.LOAD[il, :], wv @ self.PV[iv, :], il, iv

    def energy(self, n: int) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
        """返回预测电量 (ℓ̂, v̂, N̂)，单位 kWh。第 0 天无历史，返回 None。"""
        if n == 0:
            return None, None, None
        lhat, vhat, _, _ = self.raw(n)
        l_kwh, v_kwh = lhat * TAU, vhat * TAU
        return l_kwh, v_kwh, l_kwh - v_kwh


def build_error_table(LOAD: np.ndarray, PV: np.ndarray,
                      fo: Forecaster) -> np.ndarray:
    """预计算全部日期的净负荷预测误差 ε_{n,t} = N_{n,t} - N̂_{n,t|日期<n}。

    误差与 α、ρ、λ 无关，故只需算一次。第 0 天无预测，记为 NaN 并永不使用。
    这是"逐日预测后得到的误差"，不是模型在训练数据上的拟合残差。
    """
    eps = np.full((N_DAY, N), np.nan)
    for n in range(1, N_DAY):
        _, _, nhat = fo.energy(n)
        eps[n, :] = (LOAD[n, :] - PV[n, :]) * TAU - nhat
    return eps


@dataclass(frozen=True)
class RiskParams:
    """风险余量与实时执行参数（框架文档 3.3、5.2 节）。"""

    alpha: float | None = 0.80   # 净负荷误差的经验分位数；None = 不加风险余量
    rho: float = 1.0             # 储备线系数 R_t = E_min + rho(Ē_t - E_min)
    lam: float = LAMBDA_DEFAULT  # 日末储备不足的规划代价 元/kWh
    window: int = 28             # 误差回看窗口天数
    min_samples: int = 10        # 每时段最少样本数，不足则合并相邻时段

    def label(self) -> str:
        a = "无" if self.alpha is None else f"{self.alpha:.2f}"
        return f"α={a}, ρ={self.rho:.1f}, λ={self.lam:.2f}"


def risk_margin(eps: np.ndarray, n: int, rp: RiskParams) -> np.ndarray:
    """第 n 天各时段的风险余量 Q_{α,n,t}（kWh），只用 n 之前已结束日期的误差。"""
    if rp.alpha is None or n == 0:
        return np.zeros(N)
    lo = max(0, n - rp.window)
    hist = eps[lo:n, :]
    hist = hist[~np.isnan(hist).all(axis=1)]
    k = hist.shape[0]
    if k >= rp.min_samples:
        return np.quantile(hist, rp.alpha, axis=0)
    if k >= 3:
        # 历史不足时合并相邻时段（±2 个 10 分钟），仍只用过去数据
        pooled = np.empty((k, N))
        for t in range(N):
            a, b = max(0, t - 2), min(N, t + 3)
            pooled[:, t] = hist[:, a:b].mean(axis=1)
        return np.quantile(pooled, rp.alpha, axis=0)
    return np.zeros(N)


# ================================================================ 计划层
@dataclass
class DayAheadPlan:
    g: np.ndarray      # 计划购电量 kWh
    cb: np.ndarray     # 参考充电量 kWh
    db: np.ndarray     # 参考放电量 kWh
    wb: np.ndarray     # 参考弃用电量 kWh
    Ebar: np.ndarray   # 参考时段末储电量 kWh（144 个点）
    xi: float          # 日末储备不足量 kWh
    obj: float         # 规划目标值（含储备罚，不计入实际账单）


def _precomp() -> dict:
    """装配与日期无关的 MILP 结构，逐日只更换右端项，显著加速。"""
    nvar = 6 * N + 1                       # g|c|d|w|E|z|xi
    iG, iC, iD, iW, iE, iZ, iXi = 0, N, 2 * N, 3 * N, 4 * N, 5 * N, 6 * N

    rows, cols, vals = [], [], []

    def add(r, c, v):
        rows.append(r); cols.append(c); vals.append(v)

    # (1) 风险净负荷平衡：g_t + d_t - c_t - w_t = Ñ_t         行 0..N-1
    for i in range(N):
        add(i, iG + i, 1.0)
        add(i, iD + i, 1.0)
        add(i, iC + i, -1.0)
        add(i, iW + i, -1.0)

    # (2) 状态递推：E_t - E_{t-1} - η_c c_t + d_t/η_d = 0     行 N..2N-1
    for i in range(N):
        r = N + i
        add(r, iE + i, 1.0)
        if i > 0:
            add(r, iE + i - 1, -1.0)
        add(r, iC + i, -ETA)
        add(r, iD + i, 1.0 / ETA)

    A_eq = csr_matrix((vals, (rows, cols)), shape=(2 * N, nvar))

    # 不等式：充放电互斥 2N 条 + 日末储备 1 条
    urows, ucols, uvals = [], [], []

    def uadd(r, c, v):
        urows.append(r); ucols.append(c); uvals.append(v)

    for i in range(N):
        uadd(2 * i, iC + i, 1.0)
        uadd(2 * i, iZ + i, -M_ENERGY)       # c_i - M z_i <= 0
        uadd(2 * i + 1, iD + i, 1.0)
        uadd(2 * i + 1, iZ + i, M_ENERGY)    # d_i + M z_i <= M
    uadd(2 * N, iXi, -1.0)
    uadd(2 * N, iE + N - 1, -1.0)            # -xi - E_{N-1} <= -E_tar

    A_ub = csr_matrix((uvals, (urows, ucols)), shape=(2 * N + 1, nvar))

    lb = np.concatenate([np.zeros(N), np.zeros(N), np.zeros(N), np.zeros(N),
                         np.full(N, E_MIN), np.zeros(N), [0.0]])
    ub = np.concatenate([np.full(N, np.inf), np.full(N, M_ENERGY),
                         np.full(N, M_ENERGY), np.full(N, np.inf),
                         np.full(N, E_MAX), np.ones(N), [np.inf]])
    integ = np.zeros(nvar)
    integ[iZ:iZ + N] = 1

    return {"nvar": nvar, "A_eq": A_eq, "A_ub": A_ub, "lb": lb, "ub": ub,
            "integ": integ, "iG": iG, "iC": iC, "iD": iD, "iW": iW,
            "iE": iE, "iXi": iXi}


_PC = _precomp()


def solve_dayahead(ntilde: np.ndarray, e0: float, price: np.ndarray,
                   rp: RiskParams, etar: float = E_TAR) -> DayAheadPlan:
    """求解某一天的日前 MILP：在风险净负荷下优化购电与参考储能轨迹。

    min Σ p_t g_t + λ ξ,  ξ >= E_tar - Ē_144, ξ >= 0
    """
    # e0 只作为等式右端常数进入模型（不参与 Bounds），故必须在此校验：
    # 越界的 e0 要么被静默接受（过大），要么让首时段 Ē_{n,1} >= E_MIN 不可达
    # 而抛出难解的 infeasible（过小）。真实回放中 e0 恒为前一日末储电量，
    # 已在 [E_MIN, E_MAX] 内，正常情况下本检查不会触发。
    if not (E_MIN - 1e-9 <= e0 <= E_MAX + 1e-9):
        raise ValueError(
            f"日前计划的日初储电量 e0={e0:.6f} 超出 [{E_MIN}, {E_MAX}] kWh")

    nvar, iG, iC, iD, iW, iE, iXi = (
        _PC["nvar"], _PC["iG"], _PC["iC"], _PC["iD"], _PC["iW"],
        _PC["iE"], _PC["iXi"])

    b_eq = np.empty(2 * N)
    b_eq[:N] = ntilde                       # Ñ_t
    b_eq[N] = e0                            # E_0 = 当日实际日初储电量
    b_eq[N + 1:] = 0.0

    b_ub = np.empty(2 * N + 1)
    b_ub[:2 * N] = np.tile([0.0, M_ENERGY], N)
    b_ub[2 * N] = -etar

    cvec = np.zeros(nvar)
    cvec[iG:iG + N] = price
    cvec[iXi] = rp.lam

    res = milp(cvec,
               constraints=[LinearConstraint(_PC["A_eq"], b_eq, b_eq),
                            LinearConstraint(_PC["A_ub"], -np.inf, b_ub)],
               integrality=_PC["integ"], bounds=Bounds(_PC["lb"], _PC["ub"]))
    if not res.success:
        raise RuntimeError(f"日前 MILP 求解失败：{res.message}")

    x = res.x
    return DayAheadPlan(g=x[iG:iG + N], cb=x[iC:iC + N], db=x[iD:iD + N],
                        wb=x[iW:iW + N], Ebar=x[iE:iE + N],
                        xi=float(x[iXi]), obj=float(res.fun))


# ================================================================ 执行层
@dataclass
class DayExec:
    g: np.ndarray      # 计划购电量 kWh（0:00 冻结）
    c: np.ndarray      # 实际充电量 kWh
    d: np.ndarray      # 实际放电量 kWh
    r: np.ndarray      # 紧急购电量 kWh
    w: np.ndarray      # 实际弃用电量 kWh
    E: np.ndarray      # 实际时段末储电量 kWh
    E0: float
    cost_plan: float
    cost_emg: float

    @property
    def E_end(self) -> float:
        return float(self.E[-1])

    @property
    def cost_total(self) -> float:
        return self.cost_plan + self.cost_emg


def execute_day(plan_g: np.ndarray, Ebar: np.ndarray, l_act: np.ndarray,
                v_act: np.ndarray, e0: float, price: np.ndarray,
                rp: RiskParams, rule: str = "greedy",
                allow_storage: bool = True) -> DayExec:
    """实时执行：购电计划固定，储能按当前已观测情况因果响应。

    b_t = g_t + v_t - ℓ_t 为使用储能前的电量盈余（kWh）。
    b>0 用盈余充电；b<0 在储备线之上放电，剩余缺口紧急购电。
    allow_storage=False 时储能全程待机（用于构造无储能对照）。
    """
    c = np.zeros(N)
    d = np.zeros(N)
    r = np.zeros(N)
    w = np.zeros(N)
    E = np.zeros(N)
    Ep = float(e0)

    if not allow_storage:
        b_all = plan_g + v_act - l_act
        r[:] = np.maximum(0.0, -b_all)
        w[:] = np.maximum(0.0, b_all)
        E[:] = Ep
        return DayExec(g=plan_g.copy(), c=c, d=d, r=r, w=w, E=E, E0=float(e0),
                       cost_plan=float(np.sum(price * plan_g)),
                       cost_emg=float(np.sum(5.0 * price * r)))

    for t in range(N):
        b = plan_g[t] + v_act[t] - l_act[t]
        R = E_MIN + rp.rho * (Ebar[t] - E_MIN)     # 当天 0:00 即可算出的储备线

        if b > 0:
            bpos = min(b, M_ENERGY)
            if rule == "greedy":
                # 盈余全部用于充电，充不进去的弃用
                c[t] = min(bpos, (E_MAX - Ep) / ETA)
            elif rule == "refcap":
                # 先用任意来源的电量充到参考轨迹，超出部分只用光伏盈余充
                head = max(0.0, (Ebar[t] - Ep)) / ETA
                c_ref = min(bpos, head)
                pv_excess = max(0.0, v_act[t] - l_act[t])
                room = max(0.0, (E_MAX - Ep - ETA * c_ref)) / ETA
                c_extra = min(max(0.0, bpos - c_ref), max(0.0, M_ENERGY - c_ref),
                              room, pv_excess)
                c[t] = c_ref + c_extra
            else:
                raise ValueError(f"未知充电规则 {rule!r}")
            w[t] = b - c[t]
        else:
            # 电量不得低于储备线，故可放电量受储备线约束
            head = max(0.0, Ep - R)
            d[t] = min(-b, M_ENERGY, ETA * head)
            r[t] = max(0.0, -b - d[t])

        Ep = Ep + ETA * c[t] - d[t] / ETA
        E[t] = Ep

    cost_plan = float(np.sum(price * plan_g))
    cost_emg = float(np.sum(5.0 * price * r))
    return DayExec(g=plan_g.copy(), c=c, d=d, r=r, w=w, E=E, E0=float(e0),
                   cost_plan=cost_plan, cost_emg=cost_emg)


def cold_start_day(l_act: np.ndarray, v_act: np.ndarray,
                   price: np.ndarray) -> DayExec:
    """冷启动（第 1 天）：无历史可用，计划购电为零、储能待机、缺口紧急补购。"""
    g = np.zeros(N)
    r = np.maximum(0.0, l_act - v_act)
    w = np.maximum(0.0, v_act - l_act)
    return DayExec(g=g, c=np.zeros(N), d=np.zeros(N), r=r, w=w,
                   E=np.full(N, E_INIT), E0=E_INIT,
                   cost_plan=0.0, cost_emg=float(np.sum(5.0 * price * r)))


# ================================================================ 调度器
@dataclass
class DayRecord:
    n: int
    date: str
    l_hat: np.ndarray
    v_hat: np.ndarray
    n_hat: np.ndarray
    margin: np.ndarray
    n_risk: np.ndarray
    l_act: np.ndarray
    v_act: np.ndarray
    exec: DayExec
    Ebar: np.ndarray
    params: RiskParams
    rule: str
    price_1d: np.ndarray | None = None


class Simulator:
    """逐日推进的因果仿真器：维护实际电池状态与预测误差历史。"""

    def __init__(self, LOAD: np.ndarray, PV: np.ndarray, price: np.ndarray,
                 dates: list[str], eps: np.ndarray, fo: Forecaster):
        self.LOAD, self.PV, self.price = LOAD, PV, price
        self.dates, self.eps, self.fo = dates, eps, fo

    def run_day(self, n: int, rp: RiskParams, e0: float, rule: str = "greedy",
                allow_storage: bool = True) -> DayRecord:
        l_act = self.LOAD[n, :] * TAU
        v_act = self.PV[n, :] * TAU
        price_d = np.asarray(
            self.price[n] if np.ndim(self.price) == 2 else self.price,
            dtype=float,
        ).copy()

        if n == 0:
            ex = cold_start_day(l_act, v_act, price_d)
            z = np.zeros(N)
            return DayRecord(n=n, date=self.dates[n], l_hat=z, v_hat=z,
                             n_hat=z, margin=z, n_risk=z, l_act=l_act,
                             v_act=v_act, exec=ex, Ebar=np.full(N, E_INIT),
                             params=rp, rule=rule, price_1d=price_d)

        l_hat, v_hat, n_hat = self.fo.energy(n)
        q = risk_margin(self.eps, n, rp)
        n_risk = n_hat + q

        if allow_storage:
            plan = solve_dayahead(n_risk, e0, price_d, rp)
            g, Ebar = plan.g, plan.Ebar
        else:
            # 无储能时日前问题退化为逐段独立的新报童问题：g_t = max(Ñ_t, 0)
            g, Ebar = np.maximum(n_risk, 0.0), np.full(N, e0)

        ex = execute_day(g, Ebar, l_act, v_act, e0, price_d, rp, rule,
                         allow_storage)
        return DayRecord(n=n, date=self.dates[n], l_hat=l_hat, v_hat=v_hat,
                         n_hat=n_hat, margin=q, n_risk=n_risk, l_act=l_act,
                         v_act=v_act, exec=ex, Ebar=Ebar,
                         params=rp, rule=rule, price_1d=price_d)

    def replay(self, n0: int, n1: int, rp: RiskParams, e0: float,
               rule: str = "greedy",
               allow_storage: bool = True) -> tuple[float, float, list[DayRecord]]:
        """从 e0 出发回放 [n0, n1) 天，返回 (总费用, 日末储电量, 逐日记录)。"""
        recs = []
        e = e0
        tot = 0.0
        for n in range(n0, n1):
            rec = self.run_day(n, rp, e, rule, allow_storage)
            e = rec.exec.E_end
            tot += rec.exec.cost_total
            recs.append(rec)
        return tot, e, recs


# ================================================================ 参数标定
@dataclass(frozen=True)
class CalConfig:
    """参数网格与滚动标定设置（框架文档 3.4、5.2、7.1 节）。"""

    alphas: tuple = (0.50, 0.60, 0.70, 0.80, 0.90)
    rhos: tuple = (0.0, 1.0)
    lams: tuple = (0.0, 0.60)
    window: int = 35           # 标定回放窗口天数
    every: int = 42            # 每隔多少天重新标定一次
    start: int = REPORT_START  # 首次标定发生在正式区间的第一天
    rule: str = "greedy"
    allow_storage: bool = True
    report_start: int = REPORT_START
    report_end: int = REPORT_END

    def grid(self) -> list[RiskParams]:
        return [RiskParams(alpha=a, rho=r, lam=l)
                for a in self.alphas for r in self.rhos for l in self.lams]


def calibrate(sim: Simulator, n0: int, e_at_window_start: float,
              cal: CalConfig) -> tuple[RiskParams, list[dict]]:
    """在第 n0 天 0:00 用**此前**数据选择参数。

    做法：对网格中每个参数组合，从窗口起点开始回放窗口内的历史日期，
    比较包含五倍紧急购电费的**实际总费用**，取最小者。
    窗口内的初始储电量取主运行在**窗口起点日**（即 n0 - window 日）0:00 的
    实际值，因此回放是可行的；注意不是第 n0 天的值——第 n0 天在窗口末尾，
    用它的储电量作为窗口初值会引入未来信息。
    """
    lo = max(0, n0 - cal.window)
    rows = []
    best, best_cost = None, np.inf
    for rp in cal.grid():
        cost, _, _ = sim.replay(lo, n0, rp, e_at_window_start, cal.rule,
                                cal.allow_storage)
        rows.append({"标定日": sim.dates[n0], "窗口起": sim.dates[lo],
                     "参数": rp.label(), "窗口实际总费用_元": cost})
        if cost < best_cost - 1e-9:
            best_cost, best = cost, rp
    assert best is not None
    for row in rows:
        row["选中"] = (row["参数"] == best.label())
    return best, rows


def run_strategy(sim: Simulator, cal: CalConfig,
                 verbose: bool = True) -> tuple[list[DayRecord], list[dict]]:
    """从 2025-01-01 的 6000 kWh 出发跑完整年，期间滚动标定参数。

    1 月为训练与状态预热期；参数在正式区间内每隔 cal.every 天重新标定一次，
    每次只用该日之前的已完成历史，完全满足信息时点要求。
    """
    records: list[DayRecord] = []
    cal_rows: list[dict] = []
    e = E_INIT
    e_at_day_start = {0: E_INIT}
    params = RiskParams()

    for n in range(N_DAY):
        if n >= cal.start and (n - cal.start) % cal.every == 0:
            params, rows = calibrate(
                sim, n, e_at_day_start[max(0, n - cal.window)], cal)
            cal_rows.extend(rows)
            if verbose:
                print(f"    [{sim.dates[n]}] 选中 {params.label()}")
        rec = sim.run_day(n, params, e, cal.rule, cal.allow_storage)
        records.append(rec)
        e = rec.exec.E_end
        e_at_day_start[n + 1] = e
    return records, cal_rows


def strategy_totals(records: list[DayRecord], lo: int, hi: int) -> dict:
    """按正式区间汇总某个策略的费用与电量。"""
    rep = [r for r in records if lo <= r.n < hi]
    cp = sum(r.exec.cost_plan for r in rep)
    ce = sum(r.exec.cost_emg for r in rep)
    return {
        "天数": len(rep),
        "计划购电量_kWh": sum(float(r.exec.g.sum()) for r in rep),
        "紧急购电量_kWh": sum(float(r.exec.r.sum()) for r in rep),
        "弃电量_kWh": sum(float(r.exec.w.sum()) for r in rep),
        "计划购电费_元": cp,
        "紧急购电费_元": ce,
        "合计购电费_元": cp + ce,
        "期末储电量_kWh": float(rep[-1].exec.E_end),
        "期初储电量_kWh": float(rep[0].exec.E0),
    }


# ================================================================ 校验
CHECK_KEYS = ("物理可行性", "状态累计", "费用一致性")


def validate_day(rec: DayRecord, tol: float = 1e-6) -> dict[str, list[str]]:
    """逐日复核实际运行结果，按框架文档 7.2 节的检查项**分组**返回问题清单。

    返回 {检查项: [问题描述, ...]}，每项为空列表即该项通过。分组返回的目的
    是让「物理可行性」「状态累计」「费用一致性」成为三个可独立判定的结果，
    而不是共用一个「今天有没有任何问题」的布尔量——否则表面上是六项检查，
    实际只有四项是独立的。
    """
    phys: list[str] = []
    acc: list[str] = []
    cost: list[str] = []
    ex, l, v = rec.exec, rec.l_act, rec.v_act

    # 物理可行性 (1) 实际逐段供需平衡：g + r + v + d = ℓ + c + w
    resid = ex.g + ex.r + v + ex.d - l - ex.c - ex.w
    scale = max(1.0, float(l.max()))
    if np.abs(resid).max() > tol * scale:
        phys.append(f"供需平衡残差 {np.abs(resid).max():.3e} kWh")

    # 物理可行性 (2) 物理边界
    if ex.c.min() < -tol or ex.d.min() < -tol or ex.r.min() < -tol or ex.w.min() < -tol:
        phys.append("存在负的充电/放电/紧急购电/弃电")
    if ex.c.max() > M_ENERGY + tol:
        phys.append(f"充电量越限 {ex.c.max():.4f} > {M_ENERGY:.4f}")
    if ex.d.max() > M_ENERGY + tol:
        phys.append(f"放电量越限 {ex.d.max():.4f} > {M_ENERGY:.4f}")
    if ((ex.c > tol) & (ex.d > tol)).any():
        phys.append(f"存在 {int(((ex.c > tol) & (ex.d > tol)).sum())} 段同时充放电")

    # 物理可行性 (3) 储电量区间（含日初，共 145 个状态点）
    Eseq = np.concatenate([[ex.E0], ex.E])
    if Eseq.min() < E_MIN - tol:
        phys.append(f"储电量低于下界 {Eseq.min():.4f} < {E_MIN}")
    if Eseq.max() > E_MAX + tol:
        phys.append(f"储电量高于上界 {Eseq.max():.4f} > {E_MAX}")

    # 状态累计（问题二不能套用问题一的 Σd = 0.81Σc，除非首尾电量恰好相同）
    lhs = ex.E[-1] - ex.E0
    rhs = ETA * ex.c.sum() - ex.d.sum() / ETA
    if abs(lhs - rhs) > tol * max(1.0, abs(rhs)):
        acc.append(f"状态累计不符：ΔE={lhs:.6f} vs {rhs:.6f}")

    # 费用一致性：弃电不退款，储备罚不混入实际账单
    p_ref = rec.price_1d if rec.price_1d is not None else _PRICE_REF
    cp = float(np.sum(p_ref * ex.g))
    ce = float(np.sum(5.0 * p_ref * ex.r))
    if abs(cp - ex.cost_plan) > tol * max(1.0, abs(cp)):
        cost.append(f"计划费复算不符 {cp:.6f} vs {ex.cost_plan:.6f}")
    if abs(ce - ex.cost_emg) > tol * max(1.0, abs(ce)):
        cost.append(f"紧急费复算不符 {ce:.6f} vs {ex.cost_emg:.6f}")

    # 计划购电量在 0:00 后冻结（由架构保证，此处做存在性校验）
    if ex.g.shape != (N,) or ex.g.min() < -tol:
        phys.append("计划购电量非法")

    return {"物理可行性": phys, "状态累计": acc, "费用一致性": cost}


def validation_metrics(records: list[DayRecord], tol: float = 1e-6) -> dict:
    """汇总逐日校验的**数值**口径，供论文引用。

    「全部通过」这种布尔结论无法说明实际余量有多大，而判定阈值带有
    tol * max(1, max ℓ) 这样的尺度因子，直接写成「残差小于 1e-6 kWh」
    会低估真实阈值。这里把各检查项的实测最坏值原样报出，论文按实测值陈述。
    """
    m = {
        "供需平衡最大残差_kWh": 0.0,
        "供需平衡判定阈值_kWh": 0.0,
        "储电量越界最大量_kWh": 0.0,
        "充放电越限最大量_kWh": 0.0,
        "同时充放电时段数": 0,
        "状态累计最大绝对偏差_kWh": 0.0,
        "状态累计判定阈值_kWh": 0.0,
        "计划费最大复算偏差_元": 0.0,
        "紧急费最大复算偏差_元": 0.0,
    }
    for rec in records:
        ex, l, v = rec.exec, rec.l_act, rec.v_act
        resid = ex.g + ex.r + v + ex.d - l - ex.c - ex.w
        scale = max(1.0, float(l.max()))
        m["供需平衡最大残差_kWh"] = max(m["供需平衡最大残差_kWh"],
                                        float(np.abs(resid).max()))
        m["供需平衡判定阈值_kWh"] = max(m["供需平衡判定阈值_kWh"], tol * scale)
        Eseq = np.concatenate([[ex.E0], ex.E])
        m["储电量越界最大量_kWh"] = max(m["储电量越界最大量_kWh"],
                                        float(max(0.0, E_MIN - Eseq.min())),
                                        float(max(0.0, Eseq.max() - E_MAX)))
        m["充放电越限最大量_kWh"] = max(m["充放电越限最大量_kWh"],
                                        float(max(0.0, ex.c.max() - M_ENERGY)),
                                        float(max(0.0, ex.d.max() - M_ENERGY)))
        m["同时充放电时段数"] += int(((ex.c > tol) & (ex.d > tol)).sum())
        lhs = ex.E[-1] - ex.E0
        rhs = ETA * ex.c.sum() - ex.d.sum() / ETA
        m["状态累计最大绝对偏差_kWh"] = max(m["状态累计最大绝对偏差_kWh"],
                                            abs(float(lhs - rhs)))
        m["状态累计判定阈值_kWh"] = max(m["状态累计判定阈值_kWh"],
                                        tol * max(1.0, abs(float(rhs))))
        p_ref = rec.price_1d if rec.price_1d is not None else _PRICE_REF
        cp = float(np.sum(p_ref * ex.g))
        ce = float(np.sum(5.0 * p_ref * ex.r))
        m["计划费最大复算偏差_元"] = max(m["计划费最大复算偏差_元"],
                                        abs(cp - ex.cost_plan))
        m["紧急费最大复算偏差_元"] = max(m["紧急费最大复算偏差_元"],
                                        abs(ce - ex.cost_emg))
    return m


_PRICE_REF: np.ndarray | None = None     # 由 main 注入，供 validate_day 复算费用


# ================================================================ 结果整理
def emergency_events(r: np.ndarray, eps: float = EMG_EPS) -> list[dict]:
    """把同一天内连续 r_t > eps 的时段合并为一个事件（框架文档 7.4 节）。"""
    out, t = [], 0
    while t < N:
        if r[t] <= eps:
            t += 1
            continue
        a = t
        while t < N and r[t] > eps:
            t += 1
        out.append({
            "时段起": fmt_time(a * 10),
            "时段止": fmt_time(t * 10),
            "区间标签": f"{fmt_time(a * 10)}-{fmt_time(t * 10)}",
            "段数": t - a,
            "购电量_kWh": float(r[a:t].sum()),
        })
    return out


def day_summary(rec: DayRecord) -> dict:
    """单日汇总：论文表 1 / 表 2 / 表 3 与逐日指标的共同数据源。"""
    ex = rec.exec
    price = rec.price_1d if rec.price_1d is not None else _PRICE_REF
    g_kwh, c_kwh = ex.g, ex.c
    d_kwh, r_kwh, w_kwh = ex.d, ex.r, ex.w

    idx = {interval_label(i): i for i in range(N)}
    table1 = [{"时间段": lab, "计划购电量_kWh": float(g_kwh[idx[lab]])}
              for lab in WANT_LABELS]

    table2 = []
    for a, z in BLOCKS:
        table2.append({"时间段": block_label(a, z),
                       "充电量_kWh": float(c_kwh[a:z].sum()),
                       "放电量_kWh": float(d_kwh[a:z].sum())})

    ev = emergency_events(r_kwh)
    return {
        "date": rec.date,
        "table1": table1,
        "table2": table2,
        "events": ev,
        "plan_kwh": float(g_kwh.sum()),
        "emg_kwh": float(r_kwh.sum()),
        "total_kwh": float(g_kwh.sum() + r_kwh.sum()),
        "cost_plan": ex.cost_plan,
        "cost_emg": ex.cost_emg,
        "cost_total": ex.cost_total,
        "charge_kwh": float(c_kwh.sum()),
        "dischg_kwh": float(d_kwh.sum()),
        "dump_kwh": float(w_kwh.sum()),
        "E0": ex.E0,
        "E144": ex.E_end,
        "E_min": float(np.concatenate([[ex.E0], ex.E]).min()),
        "E_max": float(np.concatenate([[ex.E0], ex.E]).max()),
        "emg_segments": int((r_kwh > EMG_EPS).sum()),
        "emg_events": len(ev),
        "load_kwh": float(rec.l_act.sum()),
        "pv_kwh": float(rec.v_act.sum()),
        "params": rec.params.label(),
        "rule": rec.rule,
    }


def detail_frame(rec: DayRecord) -> pd.DataFrame:
    """逐时段明细：预测值、参考计划与实际执行值使用不同字段。"""
    E_prev = np.concatenate([[rec.exec.E0], rec.exec.E[:-1]])
    price = rec.price_1d if rec.price_1d is not None else _PRICE_REF
    return pd.DataFrame({
        "日期": rec.date,
        "时段起": [fmt_time(i * 10) for i in range(N)],
        "时段止": [fmt_time((i + 1) * 10) for i in range(N)],
        "电价_元每kWh": price,
        "预测负载_kWh": rec.l_hat,
        "预测光伏_kWh": rec.v_hat,
        "净负荷预测_kWh": rec.n_hat,
        "风险余量_kWh": rec.margin,
        "风险净负荷_kWh": rec.n_risk,
        "参考储电量_kWh": rec.Ebar,
        "实际负载_kWh": rec.l_act,
        "实际光伏_kWh": rec.v_act,
        "计划购电_kWh": rec.exec.g,
        "实际充电_kWh": rec.exec.c,
        "实际放电_kWh": rec.exec.d,
        "紧急购电_kWh": rec.exec.r,
        "弃电_kWh": rec.exec.w,
        "期初储电量_kWh": E_prev,
        "期末储电量_kWh": rec.exec.E,
        "计划费_元": price * rec.exec.g,
        "紧急费_元": 5.0 * price * rec.exec.r,
    })


def write_result2(summaries: list[dict], dates: list[str], path: Path) -> None:
    """按附件 5 模板写出 result2.xlsx。

    模板「计划购电量」的时间标签与 result1 一样整体错位一格（首列
    '0:10-0:20'、末列 '0:00-0:10+1'），这里按正文口径统一改写为
    '0:00-0:10' … '23:50-0:00+1'，行序与模板一一对应，数值不会错位。
    「充放电量」「紧急购电量」按模板结构展开到全部 334 天。
    另加一张「全天汇总」表，明列计划/紧急/合计费用，不改变原模板列。
    """
    import openpyxl

    wb = openpyxl.load_workbook(path)

    # ---- 计划购电量
    ws1 = wb["计划购电量"]
    ws1.cell(row=1, column=1, value="日期\\时间")
    for i in range(N):
        ws1.cell(row=1, column=2 + i, value=interval_label(i))
    ws1.cell(row=1, column=2 + N, value="全天购电量")
    ws1.cell(row=1, column=3 + N, value="全天购电费")
    for k, s in enumerate(summaries):
        r = 2 + k
        ws1.cell(row=r, column=1, value=dates[REPORT_START + k])
        for i, row in enumerate(s["table1"]):
            pass
        for i in range(N):
            ws1.cell(row=r, column=2 + i, value=round(s["plan_series"][i], 6))
        ws1.cell(row=r, column=2 + N, value=round(s["plan_kwh"], 6))
        ws1.cell(row=r, column=3 + N, value=round(s["cost_plan"], 6))

    # ---- 充放电量
    ws2 = wb["充放电量"]
    for row in range(ws2.max_row, 1, -1):
        ws2.delete_rows(row)
    ws2.cell(row=1, column=1, value="日期")
    ws2.cell(row=1, column=2, value="时间段")
    ws2.cell(row=1, column=3, value="充电量")
    ws2.cell(row=1, column=4, value="放电量")
    ws2.cell(row=1, column=5, value="时刻")
    ws2.cell(row=1, column=6, value="储电量")
    r = 2
    for k, s in enumerate(summaries):
        for j, blk in enumerate(s["table2"]):
            ws2.cell(row=r, column=1, value=dates[REPORT_START + k] if j == 0 else None)
            ws2.cell(row=r, column=2, value=blk["时间段"])
            ws2.cell(row=r, column=3, value=round(blk["充电量_kWh"], 6))
            ws2.cell(row=r, column=4, value=round(blk["放电量_kWh"], 6))
            r += 1
        ws2.cell(row=r - 6, column=5, value="0:00")
        ws2.cell(row=r - 6, column=6, value=round(s["E0"], 6))
        ws2.cell(row=r - 5, column=5, value="24:00")
        ws2.cell(row=r - 5, column=6, value=round(s["E144"], 6))

    # ---- 紧急购电量
    ws3 = wb["紧急购电量"]
    for row in range(ws3.max_row, 1, -1):
        ws3.delete_rows(row)
    ws3.cell(row=1, column=1, value="日期")
    ws3.cell(row=1, column=2, value="购电时间段")
    ws3.cell(row=1, column=3, value="购电量")
    r = 2
    n_ev = 0
    for k, s in enumerate(summaries):
        for e in s["events"]:
            ws3.cell(row=r, column=1, value=dates[REPORT_START + k] if e is s["events"][0] else None)
            ws3.cell(row=r, column=2, value=e["区间标签"])
            ws3.cell(row=r, column=3, value=round(e["购电量_kWh"], 6))
            r += 1
            n_ev += 1

    # ---- 全天汇总（新增表，不改变原模板列）
    if "全天汇总" in wb.sheetnames:
        del wb["全天汇总"]
    ws4 = wb.create_sheet("全天汇总")
    hdr = ["日期", "计划购电量_kWh", "紧急购电量_kWh", "合计购电量_kWh",
           "计划购电费_元", "紧急购电费_元", "合计购电费_元",
           "弃电量_kWh", "0:00储电量_kWh", "24:00储电量_kWh", "参数"]
    for j, h in enumerate(hdr, start=1):
        ws4.cell(row=1, column=j, value=h)
    for k, s in enumerate(summaries):
        r = 2 + k
        vals = [dates[REPORT_START + k], s["plan_kwh"], s["emg_kwh"],
                s["total_kwh"], s["cost_plan"], s["cost_emg"], s["cost_total"],
                s["dump_kwh"], s["E0"], s["E144"], s["params"]]
        for j, val in enumerate(vals, start=1):
            ws4.cell(row=r, column=j,
                     value=round(val, 6) if isinstance(val, float) else val)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    print(f"  计划购电量：{len(summaries)} 天 × {N} 段")
    print(f"  充放电量  ：{len(summaries) * 6} 行")
    print(f"  紧急购电量：{n_ev} 个事件行")


# ================================================================ 主流程
def main() -> None:
    global _PRICE_REF
    ap = argparse.ArgumentParser(description="2026 CUMCM C题 问题二 求解")
    ap.add_argument("--no-cal", action="store_true", help="跳过参数标定，直接用默认参数")
    ap.add_argument("--quick", action="store_true", help="不写出交付文件（调试用）")
    ap.add_argument("--rule", default="greedy", choices=["greedy", "refcap"])
    args = ap.parse_args()

    t_start = time.time()
    from p1_microgrid import load_attach1
    price = load_attach1()["电价"].to_numpy(float)
    _PRICE_REF = price

    LOAD, PV, dates = load_attach2()
    print("=" * 78)
    print(f"附件 2 载入：{LOAD.shape[0]} 天 × {LOAD.shape[1]} 个 10 分钟时段")
    print(f"  小区负载 {LOAD.min():.2f} ~ {LOAD.max():.2f} kW   "
          f"光伏实际 {PV.min():.2f} ~ {PV.max():.2f} kW")
    print(f"  合计负载 {(LOAD * TAU).sum() / 1000:,.1f} MWh   "
          f"合计光伏 {(PV * TAU).sum() / 1000:,.1f} MWh")

    fo = Forecaster(LOAD, PV)
    print(f"  预测器：{fo.fp_load.label()}；{fo.fp_pv.label()}")
    print("          （第 n 天只用第 n 天之前的记录，无未来信息）")

    eps = build_error_table(LOAD, PV, fo)
    print(f"  净负荷预测误差表：{np.sum(~np.isnan(eps).all(axis=1))} 天可用")

    # 预测精度（正式区间内，逐日全天电量误差）
    _la, _le, _na = 0.0, 0.0, 0.0
    _pa, _pe = 0.0, 0.0
    for n in range(REPORT_START, REPORT_END):
        lh, vh, nh = fo.energy(n)
        _la += abs((LOAD[n] * TAU).sum() - lh.sum())
        _le += (LOAD[n] * TAU).sum()
        _pa += abs((PV[n] * TAU).sum() - vh.sum())
        _pe += (PV[n] * TAU).sum()
        _na += abs(((LOAD[n] - PV[n]) * TAU).sum() - nh.sum())
    # 注意：此处不能用 k 作计数器——后文 for k, v in strat.items() 会把 k 覆盖成
    # 字符串，导致这里的 MAE 除法在 payload 构造时炸掉。
    n_rep = REPORT_END - REPORT_START
    print(f"  正式区间预测精度：负载 MAE {_la / n_rep:,.1f} kWh/日"
          f"（{_la / _le * 100:.2f}%）；光伏 MAE {_pa / n_rep:,.1f} kWh/日"
          f"（{_pa / _pe * 100:.2f}%）；净负荷 MAE {_na / n_rep:,.1f} kWh/日")

    # 净负荷被低估 / 被高估的结构（框架文档 7.1 节要求检查「净负荷被低估的情况」）
    # 上面的 MAE 是逐日全天电量口径的绝对误差，取绝对值后无法区分方向；而本题
    # 的代价是不对称的——低估触发 5 倍紧急购电，高估只造成弃电。故另按逐时段
    # 口径统计两个方向：低估时段占比与低估电量、高估时段占比与高估电量。
    _us_slots, _us_kwh, _ov_slots, _ov_kwh = 0, 0.0, 0, 0.0
    for n in range(REPORT_START, REPORT_END):
        _, _, nh = fo.energy(n)
        d = (LOAD[n] - PV[n]) * TAU - nh
        _us_slots += int((d > 0).sum())
        _ov_slots += int((d < 0).sum())
        _us_kwh += float(d[d > 0].sum())
        _ov_kwh += float(-d[d < 0].sum())
    _slots = n_rep * N
    print(f"  净负荷低估（实际 > 点预测）：{_us_slots / _slots * 100:.1f}% 的时段，"
          f"合计 {_us_kwh / n_rep:,.1f} kWh/日"
          f"（低估时段平均 {_us_kwh / _us_slots:,.1f} kWh/段）")
    print(f"  净负荷高估（实际 < 点预测）：{_ov_slots / _slots * 100:.1f}% 的时段，"
          f"合计 {_ov_kwh / n_rep:,.1f} kWh/日")

    sim = Simulator(LOAD, PV, price, dates, eps, fo)

    # ---------- 主策略：风险修正 + 储能（含滚动标定）
    print("-" * 78)
    print("【主策略】风险修正日前计划 + 实时执行 + 储能")
    if args.no_cal:
        cal = CalConfig(rule=args.rule, start=N_DAY)      # start=N_DAY 表示不标定
    else:
        cal = CalConfig(rule=args.rule)
    t0 = time.time()
    records, cal_rows = run_strategy(sim, cal)
    cal_log = [{"标定日": r["标定日"], "选中参数": r["参数"]}
               for r in cal_rows if r["选中"]]
    print(f"    完成，用时 {time.time() - t0:.1f}s，"
          f"年末储电量 {records[-1].exec.E_end:.2f} kWh")

    params = records[-1].params

    print("-" * 78)
    print(f"逐日推进完成：{len(records)} 天，用时 {time.time() - t_start:.1f}s，"
          f"年末储电量 {records[-1].exec.E_end:.2f} kWh")

    # ---------- 信息时点检查（框架文档 7.2 节第 1 条）
    # 逐日记录实际被引用的历史日下标，核对其最大值严格小于当天。
    leak = []
    for n in range(1, len(records)):
        _, _, il, iv = fo.raw(n)
        used = np.concatenate([il, iv]) if len(il) and len(iv) else np.array([])
        if len(used) == 0 or used.max() >= n:
            leak.append(dates[n])
    print(f"信息时点检查：预测引用的历史日索引上界 < 当天 —— "
          f"{'通过' if not leak else '失败 ' + str(leak[:5])}")

    # 误差分位数同样只取 eps[lo:n]（当天之前的已结束日期），由 risk_margin 结构保证
    print("          误差分位数仅用当天之前已结束日期的误差 —— 通过")
    print("          计划购电量在 0:00 冻结，执行层不修改 g —— 通过（结构保证）")

    # ---------- 跨日连续性（框架文档 7.2 节第 3 条）
    cont_bad = [records[n].date for n in range(1, len(records))
                if abs(records[n].exec.E0 - records[n - 1].exec.E_end) > 1e-9]
    print(f"跨日连续性：次日 0:00 储电量 == 前日 24:00 实际值 —— "
          f"{'通过' if not cont_bad else '失败 ' + str(cont_bad[:5])}")

    # ---------- 逐日物理校验（按检查项分组，三项独立判定）
    all_problems: list[tuple[str, str, str]] = []   # (检查项, 日期, 问题)
    problems_by_check: dict[str, list[tuple[str, str]]] = {
        key: [] for key in CHECK_KEYS}
    for rec in records:
        for key, msgs in validate_day(rec).items():
            for msg in msgs:
                problems_by_check[key].append((rec.date, msg))
                all_problems.append((key, rec.date, msg))
    for key in CHECK_KEYS:
        n_bad = len(problems_by_check[key])
        print(f"{key}：{len(records)} 天，"
              f"{'全部通过' if n_bad == 0 else f'{n_bad} 项问题'}")
    for key, d, msg in all_problems[:20]:
        print(f"   !! [{key}] {d}  {msg}")

    # 实测数值口径（论文按实测值陈述，而非只写「通过」）
    vm = validation_metrics(records)
    print(f"校验指标：供需平衡最大残差 {vm['供需平衡最大残差_kWh']:.3e} kWh"
          f"（判定阈值 {vm['供需平衡判定阈值_kWh']:.3e} kWh）；"
          f"储电量越界最大量 {vm['储电量越界最大量_kWh']:.3e} kWh；"
          f"状态累计最大偏差 {vm['状态累计最大绝对偏差_kWh']:.3e} kWh；"
          f"计划费/紧急费最大复算偏差 "
          f"{vm['计划费最大复算偏差_元']:.3e}/{vm['紧急费最大复算偏差_元']:.3e} 元")

    # 风险余量的符号结构：Q_α 是历史预测误差的经验 α 分位数，若某时段的误差分布
    # 整体偏负（长期高估），该时段的分位数也会是负的。论文不宜简单称其「正余量」，
    # 故此处把负值时段的占比与幅度如实统计出来。
    _rep_recs = [r for r in records if REPORT_START <= r.n < REPORT_END]
    _mg = np.concatenate([r.margin for r in _rep_recs])
    _mg_neg = _mg < 0
    print(f"风险余量：{_mg_neg.sum()} / {_mg.size} 个时段为负"
          f"（{_mg_neg.mean() * 100:.1f}%），"
          f"负值最小 {_mg.min():,.1f} kWh；正值最大 {_mg.max():,.1f} kWh")

    # ---------- 紧急事件合并一致性（框架文档 7.2 节第 6 条）
    ev_bad = 0
    for rec in records:
        evs = emergency_events(rec.exec.r)
        if abs(sum(e["购电量_kWh"] for e in evs) - float(rec.exec.r.sum())) > 1e-6:
            ev_bad += 1
    print(f"紧急事件合并一致性：合并前后电量相等 —— "
          f"{'通过' if ev_bad == 0 else f'失败 {ev_bad} 天'}")

    # ---------- 正式区间汇总
    rep = [r for r in records if REPORT_START <= r.n < REPORT_END]
    print("-" * 78)
    print(f"正式区间 {dates[REPORT_START]} ~ {dates[REPORT_END - 1]}，"
          f"共 {len(rep)} 天")

    summaries = []
    for rec in rep:
        s = day_summary(rec)
        s["plan_series"] = rec.exec.g.copy()
        summaries.append(s)

    plan_kwh = sum(s["plan_kwh"] for s in summaries)
    emg_kwh = sum(s["emg_kwh"] for s in summaries)
    cp = sum(s["cost_plan"] for s in summaries)
    ce = sum(s["cost_emg"] for s in summaries)
    print(f"  计划购电量 {plan_kwh:,.2f} kWh   计划购电费 {cp:,.2f} 元")
    print(f"  紧急购电量 {emg_kwh:,.2f} kWh   紧急购电费 {ce:,.2f} 元")
    print(f"  合计购电量 {plan_kwh + emg_kwh:,.2f} kWh   合计费用 {cp + ce:,.2f} 元")
    print(f"  紧急费占比 {ce / (cp + ce) * 100:.2f}%   "
          f"发生紧急购电的天数 {sum(1 for s in summaries if s['emg_kwh'] > EMG_EPS)}/{len(summaries)}")
    print(f"  弃电量 {sum(s['dump_kwh'] for s in summaries):,.2f} kWh   "
          f"期末储电量 {summaries[-1]['E144']:.2f} kWh")
    print(f"  储电量区间 [{min(s['E_min'] for s in summaries):.2f}, "
          f"{max(s['E_max'] for s in summaries):.2f}] kWh")

    # ---------- 指定日期
    print("-" * 78)
    by_date = {s["date"]: s for s in summaries}
    for d in SPECIAL_DATES:
        s = by_date.get(d)
        if s is None:
            print(f"  {d}：不在正式区间内")
            continue
        print(f"  表 1（{d}）指定时段计划购电量")
        for row in s["table1"]:
            print(f"     {row['时间段']:>14}  {row['计划购电量_kWh']:10.4f} kWh")
        print(f"     {'全天计划':>12}  {s['plan_kwh']:10.4f} kWh  "
              f"计划费 {s['cost_plan']:10.4f} 元")
        print(f"     {'全天紧急':>12}  {s['emg_kwh']:10.4f} kWh  "
              f"紧急费 {s['cost_emg']:10.4f} 元")
        print(f"     {'全天合计':>12}  {s['total_kwh']:10.4f} kWh  "
              f"总费用 {s['cost_total']:10.4f} 元")
        print(f"     表 2：0:00 储电量 {s['E0']:.4f}  24:00 储电量 {s['E144']:.4f} kWh")
        if s["events"]:
            print(f"     表 3：{len(s['events'])} 个紧急购电事件")
            for e_ in s["events"]:
                print(f"        {e_['区间标签']:>18}  {e_['购电量_kWh']:10.4f} kWh")
        else:
            print("     表 3：无")

    # ---------- 对照策略（框架文档 7.1 节）
    print("-" * 78)
    print("对照策略（同一预测器与实时执行规则，各自单独标定，同一评价区间）")
    cal_start = N_DAY if args.no_cal else REPORT_START

    strat: dict[str, dict] = {}
    strat_recs: dict[str, list[DayRecord]] = {}

    k_main = "本文风险修正策略（主模型）"
    strat[k_main] = strategy_totals(records, REPORT_START, REPORT_END)
    strat_recs[k_main] = records

    if not args.no_cal:
        print("  (a) 预测均值策略：不加风险余量，ρ 与 λ 照常标定")
        recs_m, _ = run_strategy(sim, CalConfig(rule=args.rule, alphas=(None,),
                                                start=cal_start), verbose=False)
        k_mean = "预测均值策略（不加分位风险余量）"
        strat[k_mean] = strategy_totals(recs_m, REPORT_START, REPORT_END)
        strat_recs[k_mean] = recs_m

        print("  (b) 无储能：同样提前计划、同样五倍补缺，α 照常标定")
        recs_ns, _ = run_strategy(sim, CalConfig(rule=args.rule, rhos=(0.0,),
                                                 allow_storage=False,
                                                 start=cal_start), verbose=False)
        k_ns = "无储能（同样提前计划并五倍补缺）"
        strat[k_ns] = strategy_totals(recs_ns, REPORT_START, REPORT_END)
        strat_recs[k_ns] = recs_ns

    # (c) 完全预知当天实际净负荷的理想对照：同模型、同参数、同一天初储电量
    #
    # 口径说明（务必如实交代，否则会被误读成可执行方案）：该对照每天从主策略的
    # 实际日初储电量 E0 出发，目标函数里没有日末储备项（标定结果为 λ=0），也
    # 不要求日末保留电量，因此它会把电池一直放到下限才收手；而次日的 E0 又回到
    # 主策略的实际值——等于每天白拿一次「已付费的库存电量」。它不是可执行方案，
    # 只是一个乐观下界，故这里把它的日末储电量也统计出来一并报告。
    ideal = 0.0
    ideal_daily: list[float] = []
    ideal_Eend: list[float] = []
    rp_ideal = RiskParams(alpha=None, rho=params.rho, lam=params.lam)
    for rec in rep:
        plan = solve_dayahead(rec.l_act - rec.v_act, rec.exec.E0, price, rp_ideal)
        c_day = float(np.sum(price * plan.g))
        ideal += c_day
        ideal_daily.append(c_day)
        ideal_Eend.append(float(plan.Ebar[-1]))
    k_id = "事后理想（完全预知当天净负荷）"
    strat[k_id] = {
        "天数": len(rep), "计划购电量_kWh": float("nan"),
        "紧急购电量_kWh": 0.0, "弃电量_kWh": float("nan"),
        "计划购电费_元": ideal, "紧急购电费_元": 0.0, "合计购电费_元": ideal,
        "期末储电量_kWh": float("nan"), "期初储电量_kWh": float("nan"),
        "日末日均储电量_kWh": float(np.mean(ideal_Eend)),
        "日末储电量最小_kWh": float(np.min(ideal_Eend)),
        "日末储电量最大_kWh": float(np.max(ideal_Eend)),
    }
    print(f"  (c) 事后理想：合计 {ideal:,.2f} 元；日末储电量 "
          f"{np.min(ideal_Eend):,.1f}--{np.max(ideal_Eend):,.1f} kWh"
          f"（日均 {np.mean(ideal_Eend):,.1f}）")

    # 逐日费用序列（供逐月分解图使用；各策略同一评价区间、同一日期轴）
    rep_dates = [s["date"] for s in summaries]
    strat_series: dict[str, dict] = {}
    for name, recs in strat_recs.items():
        rr = [r for r in recs if REPORT_START <= r.n < REPORT_END]
        strat_series[name] = {
            "dates": [r.date for r in rr],
            "cost_plan": [r.exec.cost_plan for r in rr],
            "cost_emg": [r.exec.cost_emg for r in rr],
            "cost_total": [r.exec.cost_total for r in rr],
            "plan_kwh": [float(r.exec.g.sum()) for r in rr],
            "emg_kwh": [float(r.exec.r.sum()) for r in rr],
            "E_end": [float(r.exec.E_end) for r in rr],
        }
        if strat_series[name]["dates"] != rep_dates:
            raise RuntimeError(f"策略「{name}」的日期轴与主策略不一致")
    if k_id not in strat_series:
        strat_series[k_id] = {
            "dates": rep_dates,
            "cost_plan": ideal_daily,
            "cost_emg": [0.0] * len(rep_dates),
            "cost_total": ideal_daily,
            "plan_kwh": [float("nan")] * len(rep_dates),
            "emg_kwh": [0.0] * len(rep_dates),
            "E_end": [float("nan")] * len(rep_dates),
        }


    hdr = f"  {'策略':<26}{'总费用/元':>14}{'计划费/元':>14}{'紧急费/元':>13}{'期末储电量/kWh':>16}"
    print(hdr)
    for k, v in strat.items():
        print(f"  {k:<26}{v['合计购电费_元']:>14,.0f}{v['计划购电费_元']:>14,.0f}"
              f"{v['紧急购电费_元']:>13,.0f}{v['期末储电量_kWh']:>16,.1f}")
    if not args.no_cal:
        b0 = strat["无储能（同样提前计划并五倍补缺）"]["合计购电费_元"]
        b1 = strat["预测均值策略（不加分位风险余量）"]["合计购电费_元"]
        print(f"  → 相对「无储能」节省 {b0 - (cp + ce):,.0f} 元 "
              f"（{(b0 - (cp + ce)) / b0 * 100:.2f}%）；"
              f"相对「预测均值」节省 {b1 - (cp + ce):,.0f} 元 "
              f"（{(b1 - (cp + ce)) / b1 * 100:.2f}%）")

    if args.quick:
        print("（--quick 模式，不写出交付文件）")
        return

    # ---------- 写出交付文件
    print("-" * 78)
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if TEMPLATE2.exists() and not OUT_XLSX.exists():
        import shutil
        shutil.copyfile(TEMPLATE2, OUT_XLSX)
    write_result2(summaries, dates, OUT_XLSX)

    det = pd.concat([detail_frame(r) for r in rep], ignore_index=True)
    det.round(6).to_csv(OUT_DETAIL, index=False, encoding="utf-8-sig")

    daily = pd.DataFrame([{k: v for k, v in s.items()
                           if k not in ("table1", "table2", "events",
                                        "plan_series")}
                          for s in summaries])
    daily.round(6).to_csv(OUT_DAILY, index=False, encoding="utf-8-sig")
    if cal_rows:
        pd.DataFrame(cal_rows).round(6).to_csv(OUT_CAL, index=False,
                                               encoding="utf-8-sig")

    payload = {
        "预测器": f"{fo.fp_load.label()}；{fo.fp_pv.label()}",
        "执行规则": args.rule,
        "参数标定": cal_log,
        "reported_range": [dates[REPORT_START], dates[REPORT_END - 1]],
        "n_days": len(summaries),
        "预测精度": {
            "负载MAE_kWh每日": _la / n_rep, "负载相对误差": _la / _le,
            "光伏MAE_kWh每日": _pa / n_rep, "光伏相对误差": _pa / _pe,
            "净负荷MAE_kWh每日": _na / n_rep,
        },
        "预测误差结构": {
            "低估时段占比": _us_slots / _slots,
            "低估电量_kWh每日": _us_kwh / n_rep,
            "低估时段平均低估_kWh": _us_kwh / _us_slots,
            "高估时段占比": _ov_slots / _slots,
            "高估电量_kWh每日": _ov_kwh / n_rep,
            "风险余量负值时段占比": float(_mg_neg.mean()),
            "风险余量最小值_kWh": float(_mg.min()),
            "风险余量最大值_kWh": float(_mg.max()),
        },
        "检查": {
            "1_信息时点": len(leak) == 0,
            "2_物理可行性": len(problems_by_check["物理可行性"]) == 0,
            "3_跨日连续性": len(cont_bad) == 0,
            "4_状态累计": len(problems_by_check["状态累计"]) == 0,
            "5_费用一致性": len(problems_by_check["费用一致性"]) == 0,
            "6_汇总一致性": ev_bad == 0,
        },
        "校验指标": vm,
        "totals": {
            "计划购电量_kWh": plan_kwh, "紧急购电量_kWh": emg_kwh,
            "合计购电量_kWh": plan_kwh + emg_kwh,
            "计划购电费_元": cp, "紧急购电费_元": ce, "合计购电费_元": cp + ce,
            "紧急费占比": ce / (cp + ce),
            "弃电量_kWh": sum(s["dump_kwh"] for s in summaries),
            "期末储电量_kWh": summaries[-1]["E144"],
        },
        "对照策略": strat,
        "对照策略逐日": strat_series,
        "special_dates": {d: {k: v for k, v in by_date[d].items()
                              if k != "plan_series"}
                          for d in SPECIAL_DATES if d in by_date},
        "validation_problems": [{"检查项": k, "日期": d, "问题": m}
                                for k, d, m in all_problems],
        "report_series": {
            "dates": [s["date"] for s in summaries],
            "plan_kwh": [s["plan_kwh"] for s in summaries],
            "emg_kwh": [s["emg_kwh"] for s in summaries],
            "cost_plan": [s["cost_plan"] for s in summaries],
            "cost_emg": [s["cost_emg"] for s in summaries],
            "dump_kwh": [s["dump_kwh"] for s in summaries],
            "E0": [s["E0"] for s in summaries],
            "E144": [s["E144"] for s in summaries],
        },
    }
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2,
                                   default=float), encoding="utf-8")
    print("已写出：", OUT_XLSX.relative_to(ROOT))
    print("已写出：", OUT_DETAIL.relative_to(ROOT), f"({len(det)} 行)")
    print("已写出：", OUT_DAILY.relative_to(ROOT))
    if cal_rows:
        print("已写出：", OUT_CAL.relative_to(ROOT))
    print("已写出：", OUT_JSON.relative_to(ROOT))
    print(f"总用时 {time.time() - t_start:.1f}s")


if __name__ == "__main__":
    main()
