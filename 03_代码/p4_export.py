# -*- coding: utf-8 -*-
"""问题四（合规）：计划层用电价预测，结算用当天实际附件4。

A3：0:00 不得看见当天附件4。hat[n] 只用 price4[:n]（及附件1冷启动）。
6/12/18 调整用已实现电价对余段做偏差修正。

三臂当场跑出：
  S1  计划附件1 + 结算附件1
  S2  计划附件1 + 结算附件4（同一物理计划只换计价）
  S3  计划=因果电价预测 + 结算=当天附件4  ← 提交 result4-2/4-3
  S3* 计划=当天附件4 + 结算=当天附件4     ← 上帝视角上界，不作答

不写 result3.xlsx。

复现：
  python 03_代码/p4_export.py
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
import time
from pathlib import Path

import numpy as np
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))

from p1_microgrid import load_attach1  # noqa: E402
from p2_microgrid import (  # noqa: E402
    E_INIT,
    N,
    N_DAY,
    REPORT_END,
    REPORT_START,
    RiskParams,
    Forecaster,
    Simulator,
    build_error_table,
    day_summary,
    load_attach2,
    strategy_totals,
    validate_day,
    write_result2,
)
from p3_microgrid import ATT3, ROOT, load_all, price_row  # noqa: E402
from p3_backtest import run_backtest  # noqa: E402
from p3_export import (  # noqa: E402
    BEST,
    BLOCKS,
    SPEC,
    block_label,
    export_result3,
    fmt_slot,
    validate_day as validate_p3_day,
)

ATT4 = ROOT / "01_题目" / "C题" / "附件" / "附件4.xlsx"
TEMPLATE42 = ROOT / "01_题目" / "C题" / "附件" / "附件5" / "result4-2.xlsx"
TEMPLATE43 = ROOT / "01_题目" / "C题" / "附件" / "附件5" / "result4-3.xlsx"
OUT_DIR = ROOT / "06_支撑材料"
OUT_P42 = OUT_DIR / "result4-2.xlsx"
OUT_P43 = OUT_DIR / "result4-3.xlsx"
OUT_P3 = OUT_DIR / "result3.xlsx"
OUT_JSON = OUT_DIR / "p4_results.json"
FIG_DIR = ROOT / "04_图"

P2_S1_REGRESSION = 14021565.119404145
P3_S1_REGRESSION = 13384828.869660389

P2_CUTS = (
    (31, 0.80, 1.0, 0.00),
    (73, 0.80, 0.0, 0.00),
    (115, 0.80, 0.0, 0.00),
    (157, 0.90, 1.0, 0.00),
    (199, 0.80, 0.0, 0.00),
    (241, 0.70, 0.0, 0.00),
    (283, 0.80, 0.0, 0.00),
    (325, 0.80, 0.0, 0.00),
)


def day_price(price, n: int) -> np.ndarray:
    a = np.asarray(price, dtype=float)
    if a.ndim == 2:
        return a[n].copy()
    return a.copy()


def load_attach4() -> np.ndarray:
    df = pd.read_excel(ATT4, sheet_name=0, header=0)
    price = df.iloc[:, 1:].to_numpy(float)
    if price.shape != (365, 144):
        raise ValueError(f"附件 4 形状应为 (365, 144)，实际 {price.shape}")
    if not np.isfinite(price).all() or (price < 0).any():
        raise ValueError("附件 4 非法")
    return price


def forecast_attach4(price4: np.ndarray, price1: np.ndarray, dates,
                     window: int = 28, decay: float = 0.35) -> np.ndarray:
    """因果日前电价：hat[n] 只用 price4[:n]，第 0 天退回附件1。"""
    price4 = np.asarray(price4, dtype=float)
    price1 = np.asarray(price1, dtype=float).reshape(-1)
    n_days = price4.shape[0]
    hat = np.zeros_like(price4)
    wds = np.array([pd.Timestamp(d).dayofweek for d in dates])
    for n in range(n_days):
        idx = np.arange(max(0, n - window), n)
        if len(idx) == 0:
            hat[n] = price1
            continue
        w = decay ** ((n - 1 - idx) / 7.0)
        same = wds[idx] == wds[n]
        if same.any():
            w = w * same.astype(float)
        s = float(w.sum())
        w = (w / s) if s > 0 else np.full(len(idx), 1.0 / len(idx))
        hat[n] = w @ price4[idx]
    return hat


def bill_p2(price_1d, g, r) -> tuple[float, float, float]:
    p = np.asarray(price_1d, dtype=float)
    plan = float(np.sum(p * g))
    emg = float(np.sum(5.0 * p * r))
    return plan, emg, plan + emg


def bill_p3(price_1d, gP, gA, r) -> tuple[float, float, float]:
    p = np.asarray(price_1d, dtype=float)
    up = np.clip(gA - gP, 0.0, None)
    down = np.clip(gP - gA, 0.0, None)
    pa = float(np.sum(p * gP + 1.5 * p * up - 0.5 * p * down))
    emg = float(np.sum(5.0 * p * r))
    return pa, emg, pa + emg


def rebill_p3_day(out: dict, price_1d) -> tuple[float, float, float]:
    return bill_p3(price_1d, out["gP"], out["gA"], out["r"])


def decompose(s1: float, s2: float, s3: float) -> dict[str, float]:
    return {
        "s1": float(s1), "s2": float(s2), "s3": float(s3),
        "reprice": float(s2 - s1), "strategy": float(s3 - s2),
        "total": float(s3 - s1),
    }


def p2_params_for_day(n: int) -> RiskParams:
    rp = RiskParams()
    for start, alpha, rho, lam in P2_CUTS:
        if n >= start:
            rp = RiskParams(alpha=alpha, rho=rho, lam=lam)
    return rp


def run_p2_frozen(plan_price, LOAD, PV, dates, bill_price=None, verbose: bool = True):
    fo = Forecaster(LOAD, PV)
    eps = build_error_table(LOAD, PV, fo)
    sim = Simulator(LOAD, PV, plan_price, dates, eps, fo, bill_price=bill_price)
    records = []
    e = E_INIT
    t0 = time.time()
    for n in range(N_DAY):
        rec = sim.run_day(n, p2_params_for_day(n), e)
        records.append(rec)
        e = rec.exec.E_end
        if verbose and (n + 1) % 30 == 0:
            print(f"    p2 {dates[n]} 累计 {n + 1} 天  用时 {time.time() - t0:.1f}s",
                  flush=True)
    return records


def _p2_summaries(records):
    summaries = []
    for rec in records[REPORT_START:REPORT_END]:
        s = day_summary(rec)
        s["plan_series"] = rec.exec.g.copy()
        summaries.append(s)
    return summaries


def export_p42(records, dates) -> None:
    shutil.copyfile(TEMPLATE42, OUT_P42)
    write_result2(_p2_summaries(records), dates, OUT_P42)
    from openpyxl import load_workbook
    wb = load_workbook(OUT_P42)
    if "全天汇总" in wb.sheetnames:
        del wb["全天汇总"]
    wb.save(OUT_P42)


def _spec_from_p3(res, dates, bill_price):
    date_strs = [pd.Timestamp(d).date() for d in dates]
    spec = {}
    for s in SPEC:
        dt = pd.to_datetime(s).date()
        i = date_strs.index(dt)
        o = res[i]
        p = price_row(bill_price, i)
        want_idx = [60, 72, 84, 96, 108, 120]
        t1 = [{"slot": fmt_slot(k), "gP": float(o["gP"][k]), "gA": float(o["gA"][k])}
              for k in want_idx]
        t2 = [{"block": block_label(a, z),
               "chg": float(o["c_act"][a:z].sum()),
               "dis": float(o["d_act"][a:z].sum())}
              for a, z in BLOCKS]
        r = o["r"]
        ev, t = [], 0
        while t < N:
            if r[t] > 1e-6:
                s0 = t
                while t < N and r[t] > 1e-6:
                    t += 1
                ev.append({
                    "start": fmt_slot(s0).split("-")[0],
                    "end": fmt_slot(t - 1).split("-")[1],
                    "kwh": float(r[s0:t].sum()),
                    "cost": float(np.sum(5 * p[s0:t] * r[s0:t])),
                })
            else:
                t += 1
        spec[s] = {
            "table1": t1, "Gplan": float(o["gP"].sum()), "Gadj": float(o["gA"].sum()),
            "Gemg": float(o["r"].sum()), "cost_plan_adj": float(o["cost_plan_adj"]),
            "cost_emg": float(o["cost_emg"]), "cost_total": float(o["cost_total"]),
            "table2": t2, "E0": float(o["E0"]), "E144": float(o["E"][-1]),
            "emg_events": ev,
        }
    return spec


def _daily_from_p2(records):
    return [{"date": rec.date, "cost_plan": rec.exec.cost_plan,
             "cost_emg": rec.exec.cost_emg, "cost_total": rec.exec.cost_total,
             "g": rec.exec.g.tolist(), "r": rec.exec.r.tolist()}
            for rec in records[REPORT_START:REPORT_END]]


def _daily_from_p3(res, dates):
    rows = []
    for d in range(REPORT_START, REPORT_END):
        o = res[d]
        rows.append({
            "date": str(pd.Timestamp(dates[d]).date()),
            "cost_plan_adj": float(o["cost_plan_adj"]),
            "cost_emg": float(o["cost_emg"]), "cost_total": float(o["cost_total"]),
            "gP": o["gP"].tolist(), "gA": o["gA"].tolist(), "r": o["r"].tolist(),
        })
    return rows


def rebill_p2_records(records, price) -> dict:
    plan = emg = 0.0
    daily = []
    for rec in records[REPORT_START:REPORT_END]:
        p = day_price(price, rec.n)
        cp, ce, tot = bill_p2(p, rec.exec.g, rec.exec.r)
        plan += cp
        emg += ce
        daily.append({"date": rec.date, "cost_plan": cp, "cost_emg": ce,
                      "cost_total": tot, "g": rec.exec.g.tolist(),
                      "r": rec.exec.r.tolist()})
    return {"plan": plan, "emg": emg, "total": plan + emg, "daily": daily}


def rebill_p3_res(res, dates, price) -> dict:
    pa = emg = 0.0
    daily = []
    for d in range(REPORT_START, REPORT_END):
        p = price_row(price, d)
        cpa, ce, tot = rebill_p3_day(res[d], p)
        pa += cpa
        emg += ce
        daily.append({
            "date": str(pd.Timestamp(dates[d]).date()),
            "cost_plan_adj": cpa, "cost_emg": ce, "cost_total": tot,
            "gP": res[d]["gP"].tolist(), "gA": res[d]["gA"].tolist(),
            "r": res[d]["r"].tolist(),
        })
    return {"plan_adj": pa, "emg": emg, "total": pa + emg, "daily": daily}


def make_figures(price1, price4, hat, dates, p2_s2, p2_s3, p3_s2, p3_s3,
                 p2_dec: dict, p3_dec: dict) -> list[str]:
    from cumcm_plot import savefig, setup_plot
    import matplotlib.pyplot as plt
    setup_plot()
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    written = []
    hours = (np.arange(N) + 0.5) / 6.0
    span = price4.max(axis=1) - price4.min(axis=1)
    peak = REPORT_START + int(np.argmax(span[REPORT_START:REPORT_END]))
    date_lab = str(pd.Timestamp(dates[peak]).date()) if not isinstance(dates[peak], str) else dates[peak]
    i = peak - REPORT_START

    fig, ax1 = plt.subplots(figsize=(8.2, 4.2))
    ax1.plot(hours, price1, color="#7f8c8d", lw=1.3, label="附件1")
    ax1.plot(hours, hat[peak], color="#8e44ad", lw=1.5, label="0点电价预测（因果）")
    ax1.plot(hours, price4[peak], color="#c0392b", lw=1.5, label="附件4 当天实际（结算）")
    ax1.set_ylabel("电价（元/kWh）")
    ax1.set_xlabel("时刻（h）")
    ax2 = ax1.twinx()
    ax2.plot(hours, np.array(p2_s3[i]["g"]), color="#2980b9", lw=1.1, alpha=0.9,
             label="问题二 S3 计划购电")
    ax2.plot(hours, np.array(p3_s3[i]["gA"]), color="#27ae60", lw=1.1, alpha=0.9,
             label="问题三 S3 调整购电")
    ax2.set_ylabel("购电量（kWh/10min）")
    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, loc="upper right", fontsize=8)
    ax1.set_title(f"图P4-1 {date_lab} 预测电价 vs 实际电价（S3 不用当天实际做计划）")
    ax1.set_xlim(0, 24)
    written.append(str(savefig("p4_fig_price_buy.png")))

    x = np.arange(1, len(p2_s2) + 1)
    fig, ax = plt.subplots(figsize=(8.2, 4.2))
    ax.plot(x, np.cumsum([r["cost_total"] for r in p2_s2]) / 1e4,
            color="#2980b9", ls="--", lw=1.5, label="问题二 S2（计划附件1）")
    ax.plot(x, np.cumsum([r["cost_total"] for r in p2_s3]) / 1e4,
            color="#2980b9", lw=1.8, label="问题二 S3（计划=电价预测）")
    ax.plot(x, np.cumsum([r["cost_total"] for r in p3_s2]) / 1e4,
            color="#27ae60", ls="--", lw=1.5, label="问题三 S2（计划附件1）")
    ax.plot(x, np.cumsum([r["cost_total"] for r in p3_s3]) / 1e4,
            color="#27ae60", lw=1.8, label="问题三 S3（计划=电价预测）")
    ax.set_xlabel("2–12 月累计天数")
    ax.set_ylabel("累计购电费（万元）")
    ax.set_title("图P4-2 同一附件4结算：沿用旧计划 vs 用电价预测重做计划")
    ax.legend(loc="upper left", fontsize=8)
    written.append(str(savefig("p4_fig_cumcost.png")))

    labs = ["问题二", "问题三"]
    xx = np.arange(len(labs))
    fig, ax = plt.subplots(figsize=(7.2, 4.2))
    ax.bar(xx - 0.18, [p2_dec["reprice"] / 1e4, p3_dec["reprice"] / 1e4], 0.36,
           color="#7f8c8d", label="S2−S1 计价")
    ax.bar(xx + 0.18, [p2_dec["strategy"] / 1e4, p3_dec["strategy"] / 1e4], 0.36,
           color="#c0392b", label="S3−S2 策略（预测计划）")
    ax.axhline(0.0, color="black", lw=0.6)
    ax.set_xticks(xx)
    ax.set_xticklabels(labs)
    ax.set_ylabel("费用变化（万元）")
    ax.set_title("图P4-3 分解：计价 vs 因果电价预测策略")
    ax.legend(fontsize=8)
    written.append(str(savefig("p4_fig_decompose.png")))
    return written


def _p2_pack(recs, lo=REPORT_START, hi=REPORT_END):
    t = strategy_totals(recs, lo, hi)
    return t["计划购电费_元"], t["紧急购电费_元"], t["合计购电费_元"]


def _p3_pack(res):
    pa = sum(res[d]["cost_plan_adj"] for d in range(31, 365))
    emg = sum(res[d]["cost_emg"] for d in range(31, 365))
    return pa, emg, pa + emg


def main(argv=None) -> None:
    ap = argparse.ArgumentParser(description="问题四合规：电价预测计划 + 实际结算")
    ap.add_argument("--skip-p2", action="store_true")
    ap.add_argument("--skip-p3", action="store_true")
    ap.add_argument("--no-xlsx", action="store_true")
    ap.add_argument("--no-fig", action="store_true")
    ap.add_argument("--no-oracle", action="store_true")
    args = ap.parse_args(argv)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    p3_before = (OUT_P3.stat().st_mtime, OUT_P3.stat().st_size) if OUT_P3.exists() else None

    price1 = load_attach1()["电价"].to_numpy(float)
    price4 = load_attach4()
    LOAD, PV, dates2 = load_attach2()
    hat = forecast_attach4(price4, price1, dates2)
    mae = float(np.abs(hat[REPORT_START:REPORT_END] - price4[REPORT_START:REPORT_END]).mean())
    print(f"附件4 min={price4.min():.4f} max={price4.max():.4f}；"
          f"因果电价预测 2-12月 MAE={mae:.4f} 元/kWh", flush=True)
    payload = {
        "attach4": {"min": float(price4.min()), "max": float(price4.max())},
        "price_forecast": {
            "method": "同星期指数加权，窗口28天，衰减0.35；第0天用附件1",
            "mae_feb_dec": mae,
            "note": "hat[n] 只用 price4[:n]，0点看不见当天实际电价",
        },
        "arms": {
            "S1": "计划附件1 + 结算附件1",
            "S2": "计划附件1 + 结算附件4（同一物理计划）",
            "S3": "计划=因果电价预测 + 结算=当天附件4（提交）",
            "S3_oracle": "计划=当天附件4 + 结算=当天附件4（上帝视角上界，不作答）",
        },
    }

    p2_s2_daily = p2_s3_daily = None
    p2_dec = None
    if not args.skip_p2:
        print("问题二 S1 计划+结算=附件1", flush=True)
        recs1 = run_p2_frozen(price1, LOAD, PV, dates2, bill_price=price1)
        p1, e1, t1 = _p2_pack(recs1)
        print(f"  S1 {t1:.2f}（回归 {P2_S1_REGRESSION:.2f}）", flush=True)
        s2 = rebill_p2_records(recs1, price4)
        print(f"  S2 计价 {s2['total']:.2f}  差 {s2['total']-t1:.2f}", flush=True)
        print("问题二 S3 计划=电价预测 结算=附件4", flush=True)
        recs3 = run_p2_frozen(hat, LOAD, PV, dates2, bill_price=price4)
        p3, e3, t3 = _p2_pack(recs3)
        bad = sum(1 for rec in recs3[REPORT_START:REPORT_END]
                  if any(validate_day(rec).values()))
        print(f"  S3 {t3:.2f}  校验坏天数 {bad}", flush=True)
        p2_dec = decompose(t1, s2["total"], t3)
        print(f"  分解 计价 {p2_dec['reprice']:.2f}  策略 {p2_dec['strategy']:.2f}",
              flush=True)
        payload["p2"] = {
            "S1": {"plan": p1, "emg": e1, "total": t1},
            "S2": {"plan": s2["plan"], "emg": s2["emg"], "total": s2["total"]},
            "S3": {"plan": p3, "emg": e3, "total": t3, "validate_bad": bad},
            "decompose": p2_dec,
        }
        if not args.no_oracle:
            print("问题二 S3* 上帝视角（不上报）", flush=True)
            recs_or = run_p2_frozen(price4, LOAD, PV, dates2, bill_price=price4)
            _, _, tor = _p2_pack(recs_or)
            payload["p2"]["S3_oracle"] = {"total": tor, "wan": tor / 1e4}
            print(f"  S3* {tor:.2f}", flush=True)
        p2_s2_daily = s2["daily"]
        p2_s3_daily = _daily_from_p2(recs3)
        if not args.no_xlsx:
            export_p42(recs3, dates2)
            print("  saved", OUT_P42, flush=True)

    p3_s2_daily = p3_s3_daily = None
    p3_dec = None
    if not args.skip_p3:
        _p1, dates3, load_kw, pv_kw, fc = load_all()
        hat3 = forecast_attach4(price4, price1, dates3)
        print("问题三 S1 计划+结算=附件1", flush=True)
        res1 = run_backtest(price1, dates3, load_kw, pv_kw, fc,
                            day_range=range(0, 365), **BEST,
                            use_adjust=True, adjust_mask=(True, True, True),
                            bill_price=price1)
        s1_pa, s1_emg, s1_tot = _p3_pack(res1)
        print(f"  S1 {s1_tot:.2f}（回归 {P3_S1_REGRESSION:.2f}）", flush=True)
        s2 = rebill_p3_res(res1, dates3, price4)
        print(f"  S2 计价 {s2['total']:.2f}  差 {s2['total']-s1_tot:.2f}", flush=True)
        print("问题三 S3 计划=电价预测 结算=附件4", flush=True)
        res3 = run_backtest(hat3, dates3, load_kw, pv_kw, fc,
                            day_range=range(0, 365), **BEST,
                            use_adjust=True, adjust_mask=(True, True, True),
                            bill_price=price4)
        s3_pa, s3_emg, s3_tot = _p3_pack(res3)
        bad = 0
        for d in range(31, 365):
            probs = validate_p3_day(d, price4, load_kw, pv_kw, res3[d])
            if probs:
                bad += 1
                print(dates3[d].date(), probs)
        print(f"  S3 {s3_tot:.2f}  校验 {bad} 天", flush=True)
        p3_dec = decompose(s1_tot, s2["total"], s3_tot)
        print(f"  分解 计价 {p3_dec['reprice']:.2f}  策略 {p3_dec['strategy']:.2f}",
              flush=True)
        payload["p3"] = {
            "S1": {"plan_adj": s1_pa, "emg": s1_emg, "total": s1_tot},
            "S2": {"plan_adj": s2["plan_adj"], "emg": s2["emg"], "total": s2["total"]},
            "S3": {"plan_adj": s3_pa, "emg": s3_emg, "total": s3_tot, "validate_bad": bad},
            "decompose": p3_dec,
        }
        if not args.no_oracle:
            print("问题三 S3* 上帝视角（不上报）", flush=True)
            res_or = run_backtest(price4, dates3, load_kw, pv_kw, fc,
                                  day_range=range(0, 365), **BEST,
                                  use_adjust=True, adjust_mask=(True, True, True),
                                  bill_price=price4)
            _, _, tor = _p3_pack(res_or)
            payload["p3"]["S3_oracle"] = {"total": tor, "wan": tor / 1e4}
            print(f"  S3* {tor:.2f}", flush=True)
        payload["spec"] = _spec_from_p3(res3, dates3, price4)
        p3_s2_daily = s2["daily"]
        p3_s3_daily = _daily_from_p3(res3, dates3)
        if not args.no_xlsx:
            shutil.copyfile(TEMPLATE43, OUT_P43)
            export_result3(price4, dates3, res3, out_xlsx=OUT_P43, template=TEMPLATE43)
            print("  saved", OUT_P43, flush=True)

    if not args.no_fig and p2_s2_daily and p3_s2_daily and p2_dec and p3_dec:
        figs = make_figures(price1, price4, hat, dates2,
                            p2_s2_daily, p2_s3_daily, p3_s2_daily, p3_s3_daily,
                            p2_dec, p3_dec)
        payload["figures"] = figs
        print("  figs", figs, flush=True)

    def _cost_only(rows, keys):
        return [{k: row[k] for k in keys} for row in rows]

    if p2_s2_daily is not None:
        pd.DataFrame(_cost_only(p2_s2_daily, ("date", "cost_plan", "cost_emg", "cost_total"))
                     ).to_csv(OUT_DIR / "p4_p2_s2_daily.csv", index=False, encoding="utf-8-sig")
        pd.DataFrame(_cost_only(p2_s3_daily, ("date", "cost_plan", "cost_emg", "cost_total"))
                     ).to_csv(OUT_DIR / "p4_p2_daily.csv", index=False, encoding="utf-8-sig")
    if p3_s2_daily is not None:
        pd.DataFrame(_cost_only(p3_s2_daily, ("date", "cost_plan_adj", "cost_emg", "cost_total"))
                     ).to_csv(OUT_DIR / "p4_p3_s2_daily.csv", index=False, encoding="utf-8-sig")
        pd.DataFrame(_cost_only(p3_s3_daily, ("date", "cost_plan_adj", "cost_emg", "cost_total"))
                     ).to_csv(OUT_DIR / "p4_p3_daily.csv", index=False, encoding="utf-8-sig")
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print("saved", OUT_JSON, flush=True)

    p3_after = (OUT_P3.stat().st_mtime, OUT_P3.stat().st_size) if OUT_P3.exists() else None
    if p3_before != p3_after:
        raise RuntimeError("result3.xlsx 被改写了，停止")
    print("result3.xlsx 未改写", flush=True)


if __name__ == "__main__":
    main()
