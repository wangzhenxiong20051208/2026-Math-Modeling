# -*- coding: utf-8 -*-
r"""
2026 年高教社杯全国大学生数学建模竞赛  C 题
微网与外部电网电力调控策略 —— 问题一

单日确定性计划购电优化模型（LP / MILP）
--------------------------------------------------------------------------
实现依据：《问题一_解题思路与实现框架》
  01_题目/C题/问题一_解题思路与实现框架.md

口径（与框架文档第二节一致）：
  * 附件 1 每一行的时刻代表**前一个** 10 分钟区间，即第 i 行（0 起）对应
    区间 [i*10min, (i+1)*10min]。例如标签 0:10 的行对应 0:00-0:10。
  * 全天 144 个时段，Δt = 1/6 h；附件 1 末行标签 0:00+1 表示当天 24:00。
  * 充放电量为**微网母线侧**口径：充入 1 kWh 电池增加 0.9 kWh；
    放出 1 kWh 电池减少 1/0.9 kWh；往返效率 0.81。
  * 只购电不售电；允许弃光；不额外添加外网购电功率上限。
  * E_0 = E_144 = 6000 kWh（日循环），储电量全程保持 1200-10800 kWh。

模型层次：
  1) LP   —— 连续松弛，给出最优费用下界；
  2) MILP —— 加入二元变量 z_t 强制充放电互斥，作为完整主模型。
     若 LP 最优解本身不含同时充放电，则它已是 MILP 的最优解，该结论由
     本脚本自动判定并输出，不依赖"电价为正"这一先验假设。
--------------------------------------------------------------------------
"""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import csr_matrix

# ---------------------------------------------------------------- 路径
ROOT = Path(__file__).resolve().parents[1]
ATTACH1 = ROOT / "01_题目" / "C题" / "附件" / "附件1.xlsx"
TEMPLATE1 = ROOT / "01_题目" / "C题" / "附件" / "附件5" / "result1.xlsx"
OUT_DIR = ROOT / "08_算法修正" / "输出"
OUT_XLSX = OUT_DIR / "result1.xlsx"
OUT_JSON = OUT_DIR / "p1_results.json"
OUT_DETAIL = OUT_DIR / "p1_detail.csv"

# ---------------------------------------------------------------- 常量
N = 144                      # 一天 144 个 10 分钟时段
TAU = 10 / 60                # 每时段小时数 = 1/6 h
E_CAP = 12000.0              # 储能额定容量 kWh
E_MIN = 1200.0               # 运行储电量下界 kWh
E_MAX = 10800.0              # 运行储电量上界 kWh
E_INIT = 6000.0              # 0:00 储电量 kWh
P_MAX = 5000.0               # 最大充放电功率 kW
M_ENERGY = P_MAX * TAU       # 单个时段最大充电/放电量 kWh = 2500/3
ETA = 0.90                   # 单程充放电效率
ETA_SYM = float(np.sqrt(ETA))  # 若 90% 指往返效率，则单程取 sqrt(0.9)


# ---------------------------------------------------------------- 时间标签
def fmt_time(minutes: int) -> str:
    """把「当天第几分钟」格式化为题目约定的时刻串。1440 -> '0:00+1'。"""
    if minutes == 1440:
        return "0:00+1"
    h, m = divmod(minutes, 60)
    return f"{h}:{m:02d}"


def interval_label(i: int) -> str:
    """第 i 个 10 分钟时段（0 起）的「起始-结束」标签。"""
    return f"{fmt_time(i * 10)}-{fmt_time((i + 1) * 10)}"


def block_label(a: int, z: int) -> str:
    """4 小时汇总时段的标签，末段按题目表 2 写作 '20:00-24:00'。"""
    start_h = a * 10 // 60
    end_h = 24 if z * 10 == 1440 else z * 10 // 60
    return f"{start_h}:00-{end_h}:00"


BLOCKS = [(0, 24), (24, 48), (48, 72), (72, 96), (96, 120), (120, 144)]


# ---------------------------------------------------------------- 数据载入
def load_attach1() -> pd.DataFrame:
    """读取附件 1，并校验时段标签的完整性与单调性。"""
    df = pd.read_excel(ATTACH1, sheet_name="Sheet1")
    df.columns = ["时间", "电价", "小区负载", "光伏发电预测功率"]
    if len(df) != N:
        raise ValueError(f"附件 1 应有 {N} 行，实际 {len(df)} 行")
    if df.isna().any().any():
        raise ValueError("附件 1 存在空值")

    # 附件 1 的「时间」列混用 datetime.time 与字符串，这里统一解析成分钟数，
    # 并校验其严格递增且恰好覆盖 {10, 20, ..., 1440}。
    def to_minutes(v) -> int:
        if hasattr(v, "hour"):
            return v.hour * 60 + v.minute
        s = str(v).strip()
        if s.endswith("+1"):
            return 1440
        h, m = s.split(":")
        return int(h) * 60 + int(m)

    mins = np.array([to_minutes(v) for v in df["时间"]])
    if not np.all(np.diff(mins) > 0):
        raise ValueError("附件 1 时间标签非严格递增，可能被重排或替换")
    if list(mins) != list(range(10, 1441, 10)):
        raise ValueError("附件 1 时间标签不是 {10,20,...,1440} 分钟，时段口径不成立")

    df.insert(0, "区间序号", np.arange(N))
    df["区间标签"] = [interval_label(i) for i in range(N)]
    return df


# ---------------------------------------------------------------- 模型口径
@dataclass
class Variant:
    """求解口径（主模型与敏感性分析共用）。"""

    name: str
    eta_charge: float = ETA       # 充电效率：充入 1 kWh，电池增加 eta_c kWh
    eta_discharge: float = ETA    # 放电效率：放出 1 kWh，电池减少 1/eta_d kWh
    e0: float = E_INIT            # 日初储电量
    free_e0: bool = False         # True 则日初自由，仅要求 E0 == E144
    binary: bool = False          # True 则启用二元变量强制充放电互斥（MILP）
    block_hours: int = 0          # >0 则充放电按该小时数的粗时段恒定决策

    def label(self) -> str:
        return self.name


@dataclass
class Solution:
    g: np.ndarray    # 购电功率 kW
    c: np.ndarray    # 充电功率 kW
    d: np.ndarray    # 放电功率 kW
    s: np.ndarray    # 弃光功率 kW
    E: np.ndarray    # 各时段末储电量 kWh（E[i] 为第 i 时段末）
    E0: float        # 0:00 储电量 kWh
    obj: float       # 全天购电费 元
    variant: str = ""
    is_milp: bool = False
    mip_gap: float | None = None
    extra: dict = field(default_factory=dict)


# ---------------------------------------------------------------- 求解
def solve(df: pd.DataFrame, var: Variant) -> Solution:
    """装配并求解问题一的 LP / MILP。"""
    price = df["电价"].to_numpy(float)                  # 元/kWh
    load = df["小区负载"].to_numpy(float)                # kW
    pv = df["光伏发电预测功率"].to_numpy(float)           # kW

    ec, ed = var.eta_charge, var.eta_discharge
    nz = N if var.binary else 0

    # 变量分块： [g | c | d | s | E | z]
    iG, iC, iD, iS, iE, iZ = 0, N, 2 * N, 3 * N, 4 * N, 5 * N
    nvar = 5 * N + nz

    def blk(start: int) -> slice:
        return slice(start, start + N)

    # ---- 目标：min Σ p_t g_t Δt （g 为功率，乘 Δt 得电量）
    cvec = np.zeros(nvar)
    cvec[blk(iG)] = price * TAU

    # ---- 等式约束
    eq_rows, eq_cols, eq_vals, eq_rhs = [], [], [], []

    def add_eq(row: int, coeffs: list[tuple[int, float]], rhs: float) -> None:
        for col, val in coeffs:
            eq_rows.append(row)
            eq_cols.append(col)
            eq_vals.append(val)
        eq_rhs.append(rhs)

    # (1) 供需平衡（等式 + 显式弃光）：
    #     g_t + (pv_t - s_t) + d_t = load_t + c_t
    #  => g_t + d_t - c_t - s_t = load_t - pv_t
    for i in range(N):
        add_eq(i, [(iG + i, 1.0), (iD + i, 1.0), (iC + i, -1.0), (iS + i, -1.0)],
               load[i] - pv[i])

    # (2) 储能状态递推：E_t = E_{t-1} + eta_c*c_t*Δt - d_t*Δt/eta_d
    for i in range(N):
        row = N + i
        coeffs = [(iE + i, 1.0)]
        if i > 0:
            coeffs.append((iE + i - 1, -1.0))
        coeffs.append((iC + i, -ec * TAU))
        coeffs.append((iD + i, TAU / ed))
        if i == 0 and var.free_e0:
            coeffs.append((nvar, -1.0))     # 追加变量 E0
        add_eq(row, coeffs, var.e0 if i == 0 and not var.free_e0 else 0.0)

    n_extra = 1 if var.free_e0 else 0
    if var.free_e0:
        cvec = np.concatenate([cvec, [0.0]])
        add_eq(2 * N, [(nvar, 1.0), (iE + N - 1, -1.0)], 0.0)   # E0 == E144
    else:
        add_eq(2 * N, [(iE + N - 1, 1.0)], var.e0)              # E144 == E0 == e0

    # (3) 粗时段：组内充放电功率相等
    if var.block_hours > 0:
        per = int(var.block_hours * 60 // 10)
        row = 2 * N + 1                # 紧接在 144 条平衡 + 144 条递推 + 1 条日循环之后
        for k in range(0, N, per):
            grp = list(range(k, min(k + per, N)))
            for i in grp[1:]:
                add_eq(row, [(iC + i, 1.0), (iC + grp[0], -1.0)], 0.0)
                row += 1
                add_eq(row, [(iD + i, 1.0), (iD + grp[0], -1.0)], 0.0)
                row += 1

    A_eq = csr_matrix((eq_vals, (eq_rows, eq_cols)), shape=(len(eq_rhs), nvar + n_extra))
    b_eq = np.array(eq_rhs, dtype=float)

    constraints = [LinearConstraint(A_eq, b_eq, b_eq)]

    # ---- 充放电互斥（仅 MILP）：c_t <= P_max*z_t, d_t <= P_max*(1-z_t)
    if var.binary:
        ub_rows, ub_cols, ub_vals, ub_rhs = [], [], [], []
        for i in range(N):
            # c_i - P_max*z_i <= 0
            ub_rows += [2 * i, 2 * i]
            ub_cols += [iC + i, iZ + i]
            ub_vals += [1.0, -P_MAX]
            ub_rhs.append(0.0)
            # d_i + P_max*z_i <= P_max
            ub_rows += [2 * i + 1, 2 * i + 1]
            ub_cols += [iD + i, iZ + i]
            ub_vals += [1.0, P_MAX]
            ub_rhs.append(P_MAX)
        A_ub = csr_matrix((ub_vals, (ub_rows, ub_cols)),
                          shape=(2 * N, nvar + n_extra))
        constraints.append(LinearConstraint(A_ub, -np.inf, np.array(ub_rhs)))

    # ---- 变量边界
    lb = np.concatenate([
        np.zeros(N),            # g >= 0（不售电）
        np.zeros(N),            # c >= 0
        np.zeros(N),            # d >= 0
        np.zeros(N),            # s >= 0
        np.full(N, E_MIN),      # E 下界
        np.zeros(nz),           # z
    ])
    ub = np.concatenate([
        np.full(N, np.inf),     # g 无上限：5000 kW 限制的是储能而非外网购电
        np.full(N, P_MAX),      # c
        np.full(N, P_MAX),      # d
        pv.copy(),              # s <= 光伏可用功率（弃光不超过发电）
        np.full(N, E_MAX),      # E 上界
        np.ones(nz),            # z
    ])
    if var.free_e0:
        lb = np.concatenate([lb, [E_MIN]])
        ub = np.concatenate([ub, [E_MAX]])

    integrality = np.zeros(nvar + n_extra)
    if var.binary:
        integrality[iZ:iZ + N] = 1

    res = milp(cvec, constraints=constraints, integrality=integrality,
               bounds=Bounds(lb, ub))
    if not res.success:
        raise RuntimeError(f"求解失败：{res.message}")

    x = res.x
    E0 = float(x[nvar]) if var.free_e0 else var.e0
    return Solution(
        g=x[blk(iG)], c=x[blk(iC)], d=x[blk(iD)], s=x[blk(iS)], E=x[blk(iE)],
        E0=E0, obj=float(res.fun), variant=var.label(), is_milp=var.binary,
        mip_gap=float(getattr(res, "mip_gap", 0.0) or 0.0),
    )


# ---------------------------------------------------------------- 校验
def validate(sol: Solution, df: pd.DataFrame, var: Variant, tol: float = 1e-6) -> list[str]:
    """独立复核最优解，返回问题清单（空列表 = 全部通过）。"""
    problems: list[str] = []
    price = df["电价"].to_numpy(float)
    load = df["小区负载"].to_numpy(float)
    pv = df["光伏发电预测功率"].to_numpy(float)
    ec, ed = var.eta_charge, var.eta_discharge

    # 1) 逐段供需平衡（等式 + 显式弃光）
    resid = sol.g + (pv - sol.s) + sol.d - load - sol.c
    if np.abs(resid).max() > tol * max(1.0, load.max()):
        problems.append(f"供需平衡残差过大：{np.abs(resid).max():.3e} kW")

    # 2) 弃光量非负且不超过光伏出力
    if sol.s.min() < -tol:
        problems.append(f"弃光量为负：{sol.s.min():.6f} kW")
    if (sol.s - pv).max() > tol:
        problems.append("弃光量超过光伏可用出力")

    # 3) 充放电功率限值
    if sol.c.max() > P_MAX + tol:
        problems.append(f"充电功率越限 {sol.c.max():.4f} > {P_MAX}")
    if sol.d.max() > P_MAX + tol:
        problems.append(f"放电功率越限 {sol.d.max():.4f} > {P_MAX}")

    # 4) 储电量上下界（序列含 E0，共 145 个状态点）
    Eseq = np.concatenate([[sol.E0], sol.E])
    if Eseq.min() < E_MIN - tol:
        problems.append(f"储电量低于下界 {Eseq.min():.4f} < {E_MIN}")
    if Eseq.max() > E_MAX + tol:
        problems.append(f"储电量高于上界 {Eseq.max():.4f} > {E_MAX}")

    # 5) 状态递推自洽
    Erec = [sol.E0]
    for i in range(N):
        Erec.append(Erec[-1] + ec * sol.c[i] * TAU - sol.d[i] * TAU / ed)
    err = np.abs(np.array(Erec[1:]) - sol.E).max()
    if err > tol:
        problems.append(f"储电量递推不自洽，最大偏差 {err:.6e} kWh")

    # 6) 日循环
    if abs(sol.E[-1] - sol.E0) > tol:
        problems.append(f"E144 != E0：{sol.E0:.4f} vs {sol.E[-1]:.4f}")

    # 7) 充放电互斥（唯一一条题目未直接给出、需自证的物理约束）
    both = (sol.c > tol) & (sol.d > tol)
    if both.any():
        problems.append(f"存在 {int(both.sum())} 个时段同时充放电")

    # 8) 日循环推论：Σd = η_c·η_d·Σc
    lhs = float((sol.d * TAU).sum())
    rhs = float(ec * ed * (sol.c * TAU).sum())
    if abs(lhs - rhs) > tol * max(1.0, lhs):
        problems.append(f"Σd != η_c·η_d·Σc：{lhs:.6f} vs {rhs:.6f}")

    # 9) 目标函数复算
    obj2 = float(np.sum(price * sol.g * TAU))
    if abs(obj2 - sol.obj) > tol * max(1.0, abs(sol.obj)):
        problems.append(f"目标函数复算不符：{sol.obj:.6f} vs {obj2:.6f}")

    return problems


# ---------------------------------------------------------------- 结果整理
def summarize(sol: Solution, df: pd.DataFrame) -> dict:
    """汇总论文表 1 / 表 2 及各类统计量。"""
    price = df["电价"].to_numpy(float)
    g_kwh = sol.g * TAU
    c_kwh = sol.c * TAU
    d_kwh = sol.d * TAU
    s_kwh = sol.s * TAU

    want = ["10:00-10:10", "12:00-12:10", "14:00-14:10",
            "16:00-16:10", "18:00-18:10", "20:00-20:10"]
    idx = {interval_label(i): i for i in range(N)}
    table1 = [{"时间段": lab, "购电量": float(g_kwh[idx[lab]])} for lab in want]

    table2 = []
    for a, z in BLOCKS:
        table2.append({
            "时间段": block_label(a, z),
            "充电量": float(c_kwh[a:z].sum()),
            "放电量": float(d_kwh[a:z].sum()),
        })

    return {
        "table1": table1,
        "table2": table2,
        "total_buy_kwh": float(g_kwh.sum()),
        "total_cost": float(sol.obj),
        "E0": float(sol.E0),
        "E144": float(sol.E[-1]),
        "E_min": float(np.concatenate([[sol.E0], sol.E]).min()),
        "E_max": float(np.concatenate([[sol.E0], sol.E]).max()),
        "pv_total_kwh": float((df["光伏发电预测功率"].to_numpy(float) * TAU).sum()),
        "load_total_kwh": float((df["小区负载"].to_numpy(float) * TAU).sum()),
        "charge_total_kwh": float(c_kwh.sum()),
        "dischg_total_kwh": float(d_kwh.sum()),
        "curtail_total_kwh": float(s_kwh.sum()),
        "avg_price": float(np.sum(price * g_kwh) / g_kwh.sum()),
    }


def detail_frame(sol: Solution, df: pd.DataFrame) -> pd.DataFrame:
    """完整的内部结果明细：论文表格与结果工作簿均由它汇总，避免口径不一致。"""
    E_prev = np.concatenate([[sol.E0], sol.E[:-1]])
    return pd.DataFrame({
        "时段起": [fmt_time(i * 10) for i in range(N)],
        "时段止": [fmt_time((i + 1) * 10) for i in range(N)],
        "区间标签": [interval_label(i) for i in range(N)],
        "电价_元每kWh": df["电价"].to_numpy(float),
        "负载电量_kWh": df["小区负载"].to_numpy(float) * TAU,
        "光伏电量_kWh": df["光伏发电预测功率"].to_numpy(float) * TAU,
        "购电量_kWh": sol.g * TAU,
        "充电量_kWh": sol.c * TAU,
        "放电量_kWh": sol.d * TAU,
        "弃光量_kWh": sol.s * TAU,
        "期初储电量_kWh": E_prev,
        "期末储电量_kWh": sol.E,
        "时段费用_元": df["电价"].to_numpy(float) * sol.g * TAU,
    })


def benchmarks(df: pd.DataFrame) -> dict:
    """基准对照：储能保持不动、缺口外购、富余光伏弃用。"""
    price = df["电价"].to_numpy(float)
    load = df["小区负载"].to_numpy(float)
    pv = df["光伏发电预测功率"].to_numpy(float)

    g_base = np.maximum(0.0, load - pv)          # 功率 kW
    buy_base = float((g_base * TAU).sum())
    cost_base = float((price * g_base * TAU).sum())

    buy_np = float((load * TAU).sum())
    cost_np = float((price * load * TAU).sum())

    return {
        "无储能_购电量": buy_base,
        "无储能_购电费": cost_base,
        "无储能_弃光电量": float((np.maximum(0.0, pv - load) * TAU).sum()),
        "无光伏无储能_购电量": buy_np,
        "无光伏无储能_购电费": cost_np,
    }


# ---------------------------------------------------------------- 导出
def write_result1(sol: Solution, df: pd.DataFrame, path: Path) -> None:
    """按附件 5 模板写出 result1.xlsx（以模板为底本，保留其工作表结构）。

    注意：模板「计划购电量」的时间标签整体错位一格（首行 '0:10-0:20'、末行
    '0:00+1-0:10+1' 已越过当天）。按框架文档第二节的口径，写出正确的
    '0:00-0:10' … '23:50-0:00+1'；行序与模板一一对应，数值不会错位。
    原模板文件保持不动。
    """
    g_kwh = sol.g * TAU
    c_kwh = sol.c * TAU
    d_kwh = sol.d * TAU

    wb = openpyxl.load_workbook(path) if path.exists() else openpyxl.Workbook()

    ws1 = wb["计划购电量"] if "计划购电量" in wb.sheetnames else wb.active
    if ws1.title != "计划购电量":
        ws1.title = "计划购电量"
    ws1["A1"], ws1["B1"] = "时间段", "购电量"
    for i in range(N):
        ws1.cell(row=i + 2, column=1, value=interval_label(i))
        ws1.cell(row=i + 2, column=2, value=round(float(g_kwh[i]), 6))

    ws2 = wb["充放电量"] if "充放电量" in wb.sheetnames else wb.create_sheet("充放电量")
    ws2["A1"], ws2["B1"], ws2["C1"], ws2["D1"], ws2["E1"] = (
        "时间段", "充电量", "放电量", "时刻", "储电量")
    for r, (a, z) in enumerate(BLOCKS, start=2):
        ws2.cell(row=r, column=1, value=block_label(a, z))
        ws2.cell(row=r, column=2, value=round(float(c_kwh[a:z].sum()), 6))
        ws2.cell(row=r, column=3, value=round(float(d_kwh[a:z].sum()), 6))
    ws2.cell(row=2, column=4, value="0:00")
    ws2.cell(row=2, column=5, value=round(float(sol.E0), 6))
    ws2.cell(row=3, column=4, value="24:00")
    ws2.cell(row=3, column=5, value=round(float(sol.E[-1]), 6))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ---------------------------------------------------------------- 主流程
def main() -> None:
    ap = argparse.ArgumentParser(description="2026 CUMCM C题 问题一 求解")
    ap.add_argument("--all-variants", action="store_true", help="同时求解敏感性分析口径")
    args = ap.parse_args()

    df = load_attach1()
    price = df["电价"].to_numpy(float)
    load = df["小区负载"].to_numpy(float)
    pv = df["光伏发电预测功率"].to_numpy(float)

    print("=" * 78)
    print(f"附件 1 载入：{len(df)} 个 10 分钟时段")
    print(f"  电价 {price.min():.4f} ~ {price.max():.4f} 元/kWh   "
          f"负载 {load.min():.2f} ~ {load.max():.2f} kW   "
          f"光伏 {pv.min():.2f} ~ {pv.max():.2f} kW")
    print(f"  时段标签校验：严格递增且恰为 {{10,20,...,1440}} 分钟 —— 通过")

    # ---------- 第 1 步：LP（连续松弛，给出费用下界）
    lp_var = Variant(name="LP（连续松弛，费用下界）")
    lp_sol = solve(df, lp_var)
    lp_problems = validate(lp_sol, df, lp_var)
    lp_both = int(((lp_sol.c > 1e-9) & (lp_sol.d > 1e-9)).sum())
    print("-" * 78)
    print(f"LP   ：最优费用 {lp_sol.obj:,.4f} 元   "
          f"校验{'通过' if not lp_problems else '存在问题'}")
    for p in lp_problems:
        print("   !!", p)
    print(f"       同时充放电的时段数 = {lp_both}")

    # ---------- 第 2 步：MILP（含二元变量，强制充放电互斥）
    milp_var = Variant(name="MILP（含充放电互斥约束，完整主模型）", binary=True)
    milp_sol = solve(df, milp_var)
    milp_problems = validate(milp_sol, df, milp_var)
    print(f"MILP ：最优费用 {milp_sol.obj:,.4f} 元   "
          f"校验{'通过' if not milp_problems else '存在问题'}   "
          f"gap={milp_sol.mip_gap:.2e}")
    for p in milp_problems:
        print("   !!", p)

    # ---------- 第 3 步：互斥性判定
    # 若 LP 最优解本身不含同时充放电，则它满足完整模型，即为 MILP 最优解。
    gap = abs(milp_sol.obj - lp_sol.obj)
    print("-" * 78)
    if lp_both == 0 and gap <= 1e-6 * max(1.0, abs(milp_sol.obj)):
        print("结论：LP 最优解不含同时充放电，且 LP 与 MILP 目标值一致")
        print("      => LP 松弛的最优解已是完整模型（MILP）的最优解，")
        print("         无需依赖『电价为正』这一先验假设。")
    else:
        print(f"结论：需采用 MILP 结果。LP 同时充放电时段数={lp_both}，"
              f"两者费用差 {gap:.6f} 元")
    sol = milp_sol if gap > 1e-6 * max(1.0, abs(milp_sol.obj)) else lp_sol

    # ---------- 主模型结果
    var = Variant(name="主模型：LP⊂MILP，η充=η放=0.90，E0=6000")
    s = summarize(sol, df)
    print("-" * 78)
    print(f"全天购电量 {s['total_buy_kwh']:,.2f} kWh   "
          f"全天购电费 {s['total_cost']:,.2f} 元   "
          f"平均购电单价 {s['avg_price']:.4f} 元/kWh")
    print(f"光伏发电 {s['pv_total_kwh']:,.2f} kWh   负载需求 {s['load_total_kwh']:,.2f} kWh "
          f"  弃光 {s['curtail_total_kwh']:,.2f} kWh")
    print(f"储能充电 {s['charge_total_kwh']:,.2f} kWh   放电 {s['dischg_total_kwh']:,.2f} kWh"
          f"   比值 {s['dischg_total_kwh'] / s['charge_total_kwh']:.6f}（应为 0.81）")
    print(f"储电量区间 [{s['E_min']:.2f}, {s['E_max']:.2f}] kWh   "
          f"0:00={s['E0']:.2f}   24:00={s['E144']:.2f}")

    print("-" * 78)
    print("表 1  指定时段的购电量")
    for row in s["table1"]:
        print(f"  {row['时间段']:>14}  {row['购电量']:12.4f} kWh")
    print(f"  {'全天购电量':>12}  {s['total_buy_kwh']:12.4f} kWh")
    print(f"  {'全天购电费':>12}  {s['total_cost']:12.4f} 元")

    print("-" * 78)
    print("表 2  储能设备充放电量")
    for row in s["table2"]:
        print(f"  {row['时间段']:>14}  充电 {row['充电量']:10.4f}   "
              f"放电 {row['放电量']:10.4f} kWh")
    print(f"  0:00 储电量 {s['E0']:.4f} kWh    24:00 储电量 {s['E144']:.4f} kWh")

    bm = benchmarks(df)
    print("-" * 78)
    print("基准对照（储能保持不动、缺口外购、富余光伏弃用）")
    print(f"  {'基准（无储能）':<16} 购电量 {bm['无储能_购电量']:12.2f} kWh   "
          f"购电费 {bm['无储能_购电费']:11.2f} 元  弃光 {bm['无储能_弃光电量']:9.2f} kWh")
    print(f"  {'含储能（本模型）':<16} 购电量 {s['total_buy_kwh']:12.2f} kWh   "
          f"购电费 {s['total_cost']:11.2f} 元  弃光 {s['curtail_total_kwh']:9.2f} kWh")
    save_amt = bm["无储能_购电费"] - s["total_cost"]
    print(f"  → 节省 {save_amt:.2f} 元，节费率 {save_amt / bm['无储能_购电费'] * 100:.2f}%"
          f"（最优费用不高于该可行基准：{s['total_cost'] <= bm['无储能_购电费']}）")

    # ---------- 写出交付文件
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    if TEMPLATE1.exists() and not OUT_XLSX.exists():
        import shutil
        shutil.copyfile(TEMPLATE1, OUT_XLSX)
    write_result1(sol, df, OUT_XLSX)
    detail_frame(sol, df).to_csv(OUT_DETAIL, index=False, encoding="utf-8-sig")
    print("-" * 78)
    print("已写出：", OUT_XLSX.relative_to(ROOT))
    print("已写出：", OUT_DETAIL.relative_to(ROOT))

    # ---------- 敏感性分析
    payload = {"main": {"variant": var.name, **s}, "benchmarks": bm}
    if args.all_variants:
        variants = [
            Variant(name="主模型：η充=η放=0.90（往返0.81），E0=6000"),
            Variant(name="放电无损：η充=0.90，η放=1.00（往返0.90）",
                    eta_charge=ETA, eta_discharge=1.0),
            Variant(name="对称折算：η充=η放=√0.90（往返0.90）",
                    eta_charge=ETA_SYM, eta_discharge=ETA_SYM),
            Variant(name="日初自由：仅要求 E0=E144", free_e0=True),
            Variant(name="充放电按 1 小时粗时段决策", block_hours=1),
            Variant(name="充放电按 4 小时粗时段决策", block_hours=4),
        ]
        payload["variants"] = []
        for v in variants:
            try:
                sv = solve(df, v)
                sm = summarize(sv, df)
                payload["variants"].append({
                    "口径": v.name,
                    "全天购电量_kWh": sm["total_buy_kwh"],
                    "全天购电费_元": sm["total_cost"],
                    "充电量_kWh": sm["charge_total_kwh"],
                    "放电量_kWh": sm["dischg_total_kwh"],
                    "E0": sm["E0"], "E144": sm["E144"],
                    "弃光_kWh": sm["curtail_total_kwh"],
                    "问题": validate(sv, df, v),
                })
            except Exception as exc:  # noqa: BLE001
                payload["variants"].append({"口径": v.name, "错误": str(exc)})
        print("-" * 78)
        print("敏感性分析")
        for row in payload["variants"]:
            if "错误" in row:
                print(f"  {row['口径']:<40}  {row['错误']}")
            else:
                print(f"  {row['口径']:<34} 购电量 {row['全天购电量_kWh']:11.2f} kWh  "
                      f"购电费 {row['全天购电费_元']:10.2f} 元  "
                      f"充 {row['充电量_kWh']:9.2f}  放 {row['放电量_kWh']:9.2f} kWh  "
                      f"放/充 {row['放电量_kWh'] / row['充电量_kWh']:.4f}")

    payload["milp_check"] = {
        "LP费用": lp_sol.obj, "MILP费用": milp_sol.obj,
        "LP同时充放电时段数": lp_both, "费用差": gap,
        "LP解即MILP最优": bool(lp_both == 0 and gap <= 1e-6 * max(1.0, abs(milp_sol.obj))),
    }
    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print("已写出：", OUT_JSON.relative_to(ROOT))


if __name__ == "__main__":
    main()
