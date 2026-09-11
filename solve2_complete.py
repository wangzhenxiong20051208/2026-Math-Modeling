# -*- coding: utf-8 -*-
"""
2026 CUMCM C题 问题二 — 完整求解脚本（一步到位输出result2.xlsx）
严格合规：
1. 预测参数在1月离线选定（只用1月数据）
2. α/ρ每42天滚动标定（用过去35天实际运行费用）
3. 全程不使用未来信息
"""
import numpy as np
import pandas as pd
from scipy.optimize import linprog
import openpyxl
import datetime

# ==================== 常量 ====================
N = 144
DT = 1 / 6
ETA = 0.9
EMIN = 1200.0
EMAX = 10800.0
E_INIT = 6000.0
PMAX = 5000.0
M = PMAX * DT
REPORT_START = 31  # 2月1日
N_DAY = 365
CAL_EVERY = 42  # 每42天标定一次
CAL_WINDOW = 35  # 标定回放窗口35天

# ==================== 步骤1: 数据加载 ====================
print("=" * 60)
print("步骤1: 加载数据")
print("=" * 60)

df1 = pd.read_excel(r'01_题目\C题\附件\附件1.xlsx')
price = df1['电价'].astype(float).values

load_df = pd.read_excel(r'01_题目\C题\附件\附件2.xlsx', sheet_name='小区负载')
pv_df = pd.read_excel(r'01_题目\C题\附件\附件2.xlsx', sheet_name='光伏发电实际功率')
dates = pd.to_datetime(load_df.iloc[:, 0]).tolist()
load_mat = load_df.iloc[:, 1:].astype(float).values
pv_mat = pv_df.iloc[:, 1:].astype(float).values
load_e = load_mat * DT
pv_e = pv_mat * DT

print(f"  电价: {price.shape}, 负载: {load_e.shape}, 光伏: {pv_e.shape}")


def weekday_of(n):
    return (n + 2) % 7


# ==================== 预测函数 ====================
def predict_day(n, load_win, load_decay, pv_win, pv_decay):
    """预测第n天的净负荷（因果，只用历史数据）"""
    if n == 0:
        return np.zeros(N)
    # 负载：同星期优先，指数衰减加权
    lo = max(0, n - load_win)
    idx = np.arange(lo, n)
    same = np.array([weekday_of(i) == weekday_of(n) for i in idx])
    idx_same = idx[same] if same.any() else idx
    w = load_decay ** ((n - 1 - idx_same) / 7.0)
    w = w / w.sum()
    pl = (w[:, None] * load_e[idx_same, :]).sum(axis=0)
    # 光伏：指数衰减加权
    lo2 = max(0, n - pv_win)
    idx2 = np.arange(lo2, n)
    w2 = pv_decay ** ((n - 1 - idx2) / 7.0)
    w2 = w2 / w2.sum()
    pp = (w2[:, None] * pv_e[idx2, :]).sum(axis=0)
    return pl - pp


# ==================== 安全裕度 ====================
def safety_margin(eps_table, n, alpha, window=28, min_samples=8):
    if n == 0:
        return np.zeros(N)
    lo = max(0, n - window)
    hist = eps_table[lo:n, :]
    valid = ~np.isnan(hist).all(axis=1)
    hist = hist[valid]
    k = hist.shape[0]
    if k >= min_samples:
        return np.nanquantile(hist, alpha, axis=0)
    if k >= 3:
        pooled = np.empty((k, N))
        for t in range(N):
            a, b = max(0, t - 2), min(N, t + 3)
            pooled[:, t] = np.nanmean(hist[:, a:b], axis=1)
        return np.nanquantile(pooled, alpha, axis=0)
    return np.zeros(N)


# ==================== 日LP（0:00计划层） ====================
def solve_plan(net_load, e0):
    iG, iC, iD, iS, iE = 0, N, 2 * N, 3 * N, 4 * N
    nv = 5 * N
    c_obj = np.zeros(nv)
    c_obj[iG:iG + N] = price
    A_eq = np.zeros((2 * N, nv))
    b_eq = np.zeros(2 * N)
    for t in range(N):
        A_eq[t, iG + t] = 1
        A_eq[t, iD + t] = 1
        A_eq[t, iC + t] = -1
        A_eq[t, iS + t] = -1
        b_eq[t] = net_load[t]
    for t in range(N):
        A_eq[N + t, iE + t] = 1
        if t > 0:
            A_eq[N + t, iE + t - 1] = -1
        A_eq[N + t, iC + t] = -ETA
        A_eq[N + t, iD + t] = 1.0 / ETA
        if t == 0:
            b_eq[N + t] = e0
    bounds = ([(0, None)] * N + [(0, M)] * N + [(0, M)] * N
              + [(0, None)] * N + [(EMIN, EMAX)] * N)
    res = linprog(c_obj, A_eq=A_eq, b_eq=b_eq, bounds=bounds, method='highs')
    if not res.success:
        return None, None
    return res.x[iG:iG + N].copy(), res.x[iE:iE + N].copy()


# ==================== 执行层（贪心+储备线） ====================
def execute_day(g_plan, E_ref, l_act, v_act, e0, rho):
    c = np.zeros(N)
    d = np.zeros(N)
    r = np.zeros(N)
    w = np.zeros(N)
    E = np.zeros(N)
    soc = e0
    for t in range(N):
        b = g_plan[t] + v_act[t] - l_act[t]
        R = EMIN + rho * (E_ref[t] - EMIN)
        if b >= 0:
            cmax = min(b, M, (EMAX - soc) / ETA)
            c[t] = max(0, cmax)
            w[t] = b - c[t]
        else:
            head = max(0.0, soc - R)
            dmax = min(-b, M, ETA * head)
            d[t] = max(0, dmax)
            r[t] = max(0, -b - d[t])
        soc = soc + ETA * c[t] - d[t] / ETA
        E[t] = soc
    return c, d, r, w, E, float(np.sum(price * g_plan)), float(np.sum(5.0 * price * r))


# ==================== 单日运行 ====================
def run_day(n, e0, alpha, rho, pred_params, eps_table):
    load_win, load_decay, pv_win, pv_decay = pred_params
    pnet = predict_day(n, load_win, load_decay, pv_win, pv_decay)
    sm = safety_margin(eps_table, n, alpha)
    ntilde = pnet + sm
    g, E_ref = solve_plan(ntilde, e0)
    if g is None:
        return None
    c, d, r, w, E, cp, ce = execute_day(g, E_ref, load_e[n], pv_e[n], e0, rho)
    return {
        'n': n, 'date': dates[n], 'g': g, 'c': c, 'd': d, 'r': r, 'w': w, 'E': E,
        'E0': e0, 'E_end': E[-1], 'cost_plan': cp, 'cost_emg': ce,
        'cost_total': cp + ce, 'alpha': alpha, 'rho': rho
    }


# ==================== 步骤2: 1月离线选定预测参数 ====================
print("\n" + "=" * 60)
print("步骤2: 1月离线选定预测参数")
print("=" * 60)

pred_candidates = []
for lw in [28, 35]:
    for ld in [0.3, 0.5]:
        for pw in [7, 14]:
            for pd in [0.3, 0.5]:
                pred_candidates.append((lw, ld, pw, pd))

best_pred_cost = np.inf
best_pred_params = None

for params in pred_candidates:
    lw, ld, pw, pd = params
    # 构建1月的误差表
    eps_jan = np.full((31, N), np.nan)
    for n in range(1, 31):
        pnet = predict_day(n, lw, ld, pw, pd)
        eps_jan[n] = (load_e[n] - pv_e[n]) - pnet
    # 跑1月
    e = E_INIT
    total_cost = 0.0
    for n in range(31):
        rec = run_day(n, e, 0.80, 0.75, params, eps_jan)
        if rec is None:
            total_cost = np.inf
            break
        total_cost += rec['cost_total']
        e = rec['E_end']
    print(f"  负载({lw},{ld}) 光伏({pw},{pd}): 1月总费={total_cost:.2f}")
    if total_cost < best_pred_cost:
        best_pred_cost = total_cost
        best_pred_params = params

lw, ld, pw, pd = best_pred_params
print(f"\n最优预测参数: 负载({lw},{ld}) 光伏({pw},{pd})")
print(f"1月总费用: {best_pred_cost:.2f}")

# ==================== 步骤3: 构建全年误差表 ====================
print("\n" + "=" * 60)
print("步骤3: 构建全年预测和误差表")
print("=" * 60)

pred_net_full = np.zeros((N_DAY, N))
eps_table_full = np.full((N_DAY, N), np.nan)
for n in range(1, N_DAY):
    pred_net_full[n] = predict_day(n, lw, ld, pw, pd)
    eps_table_full[n] = (load_e[n] - pv_e[n]) - pred_net_full[n]
print("完成")

# ==================== 步骤4: 1月预热期 ====================
print("\n" + "=" * 60)
print("步骤4: 1月预热期运行（默认α=0.80, ρ=0.75）")
print("=" * 60)

e = E_INIT
e_at_start = np.zeros(N_DAY + 1)
e_at_start[0] = E_INIT
warmup_records = []
for n in range(0, REPORT_START):
    rec = run_day(n, e, 0.80, 0.75, best_pred_params, eps_table_full)
    warmup_records.append(rec)
    e = rec['E_end']
    e_at_start[n + 1] = e
print(f"1月预热完成，2.1初始SOC={e:.2f}")

# ==================== 步骤5: 正式区间滚动标定+运行 ====================
print("\n" + "=" * 60)
print(f"步骤5: 正式区间（2.1-12.31）滚动标定+运行")
print(f"标定频率: 每{CAL_EVERY}天, 回放窗口: {CAL_WINDOW}天")
print("=" * 60)

alpha_candidates = [0.70, 0.75, 0.80, 0.85, 0.90]
rho_candidates = [0.5, 0.6, 0.7, 0.8, 0.9, 1.0]


def replay_window(n_start, n_end, e_start, alpha, rho):
    """从n_start到n_end-1回放，返回总费用"""
    e = e_start
    total = 0.0
    for n in range(n_start, n_end):
        rec = run_day(n, e, alpha, rho, best_pred_params, eps_table_full)
        if rec is None:
            return np.inf, e
        total += rec['cost_total']
        e = rec['E_end']
    return total, e


all_records = []
calibration_log = []
current_alpha = 0.80
current_rho = 0.75

for n in range(REPORT_START, N_DAY):
    # 检查是否需要标定
    if (n - REPORT_START) % CAL_EVERY == 0:
        print(f"\n  [标定] {dates[n].date()} (n={n})...")
        lo = max(0, n - CAL_WINDOW)
        e_win_start = e_at_start[lo]
        best_cost = np.inf
        best_ar = None
        for a in alpha_candidates:
            for r in rho_candidates:
                cost, _ = replay_window(lo, n, e_win_start, a, r)
                if cost < best_cost - 1e-9:
                    best_cost = cost
                    best_ar = (a, r)
        current_alpha, current_rho = best_ar
        calibration_log.append({
            'n': n, 'date': dates[n],
            'alpha': current_alpha, 'rho': current_rho
        })
        print(f"  选中: α={current_alpha}, ρ={current_rho}, 窗口费={best_cost:.2f}")

    rec = run_day(n, e, current_alpha, current_rho, best_pred_params, eps_table_full)
    all_records.append(rec)
    e = rec['E_end']
    e_at_start[n + 1] = e

    if (n - REPORT_START) % 50 == 0 or n == N_DAY - 1:
        print(f"  {dates[n].date()} E0={rec['E0']:.1f} E_end={e:.1f} "
              f"plan={rec['cost_plan']:.1f} emg={rec['cost_emg']:.1f} "
              f"α={current_alpha} ρ={current_rho}")

# ==================== 步骤6: 汇总 ====================
print("\n" + "=" * 60)
print("步骤6: 结果汇总")
print("=" * 60)

formal = all_records
total_plan = sum(r['cost_plan'] for r in formal)
total_emg = sum(r['cost_emg'] for r in formal)
total_g = sum(float(r['g'].sum()) for r in formal)
total_r = sum(float(r['r'].sum()) for r in formal)
total_w = sum(float(r['w'].sum()) for r in formal)
days_emg = sum(1 for r in formal if r['r'].sum() > 1e-6)

print(f"预测参数: 负载({lw},{ld}) 光伏({pw},{pd})")
print(f"计划购电量: {total_g:.2f} kWh")
print(f"紧急购电量: {total_r:.2f} kWh")
print(f"计划购电费: {total_plan:.2f} 元")
print(f"紧急购电费: {total_emg:.2f} 元")
print(f"合计购电费: {total_plan + total_emg:.2f} 元")
print(f"紧急费占比: {total_emg / (total_plan + total_emg) * 100:.2f}%")
print(f"弃用电量: {total_w:.2f} kWh")
print(f"期末储电量: {formal[-1]['E_end']:.2f} kWh")
print(f"发生紧急购电天数: {days_emg}/334")

print(f"\n标定日志:")
for cl in calibration_log:
    print(f"  {cl['date'].date()}: α={cl['alpha']}, ρ={cl['rho']}")

# ==================== 步骤7: 写入result2.xlsx ====================
print("\n" + "=" * 60)
print("步骤7: 写入result2.xlsx")
print("=" * 60)

wb = openpyxl.load_workbook(r'01_题目\C题\附件\附件5\result2.xlsx')
print(f"  Sheet列表: {wb.sheetnames}")

# Sheet1: 计划购电量
ws = wb['计划购电量']
if ws.max_row > 1:
    ws.delete_rows(2, ws.max_row - 1)
for i, rec in enumerate(formal):
    row = i + 2
    ws.cell(row, 1).value = rec['date']
    g = rec['g']
    for t in range(144):
        ws.cell(row, t + 2).value = round(float(g[t]), 4)
    ws.cell(row, 146).value = round(float(g.sum()), 2)
    ws.cell(row, 147).value = round(float(np.sum(price * g)), 2)
print("  计划购电量 sheet 完成")

# Sheet2: 充放电量（如果有这个sheet）
if '充放电量' in wb.sheetnames:
    ws3 = wb['充放电量']
    if ws3.max_row > 1:
        ws3.delete_rows(2, ws3.max_row - 1)
    time_slots = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00', '16:00-20:00', '20:00-24:00']
    for i, rec in enumerate(formal):
        base_row = i * 6 + 2
        c = rec['c']
        d = rec['d']
        E0 = rec['E0']
        E_end = rec['E_end']
        for k in range(6):
            row = base_row + k
            ws3.cell(row, 2).value = time_slots[k]
            c_sum = float(c[k * 24:(k + 1) * 24].sum())
            d_sum = float(d[k * 24:(k + 1) * 24].sum())
            ws3.cell(row, 3).value = round(c_sum, 4)
            ws3.cell(row, 4).value = round(d_sum, 4)
            if k == 0:
                ws3.cell(row, 1).value = rec['date']
                ws3.cell(row, 5).value = datetime.time(0, 0)
                ws3.cell(row, 6).value = round(float(E0), 2)
            elif k == 1:
                ws3.cell(row, 5).value = '24:00'
                ws3.cell(row, 6).value = round(float(E_end), 2)
    print("  充放电量 sheet 完成")

# Sheet3: 紧急购电量
if '紧急购电量' in wb.sheetnames:
    ws4 = wb['紧急购电量']
    if ws4.max_row > 1:
        ws4.delete_rows(2, ws4.max_row - 1)
    emg_events = []
    for rec in formal:
        r = rec['r']
        date = rec['date']
        t = 0
        while t < 144:
            if r[t] > 1e-6:
                start_t = t
                emg_sum = 0.0
                while t < 144 and r[t] > 1e-6:
                    emg_sum += r[t]
                    t += 1
                end_t = t - 1
                start_min = start_t * 10
                end_min = (end_t + 1) * 10
                start_label = f"{start_min // 60}:{start_min % 60:02d}"
                end_label = "0:00+1" if end_min >= 1440 else f"{end_min // 60}:{end_min % 60:02d}"
                emg_events.append((date, f"{start_label}-{end_label}", round(float(emg_sum), 2)))
            else:
                t += 1
    for i, (date, time_label, amount) in enumerate(emg_events):
        row = i + 2
        ws4.cell(row, 1).value = date
        ws4.cell(row, 2).value = time_label
        ws4.cell(row, 3).value = amount
    print(f"  紧急购电量 sheet 完成 ({len(emg_events)}个事件)")

wb.save(r'result2.xlsx')

# ==================== 完成 ====================
print("\n" + "=" * 60)
print("完成！")
print("=" * 60)
print(f"  输出文件: result2.xlsx")
print(f"  最终总费用: {total_plan + total_emg:.2f} 元")
print(f"  预测参数: 负载({lw},{ld}) 光伏({pw},{pd})")
print(f"  合规性: 1月标定预测参数 -> 滚动标定α/ρ -> 2-12月求解，无事后选参")
