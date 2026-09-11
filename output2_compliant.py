# -*- coding: utf-8 -*-
"""
将严格合规版结果写入 result2_compliant.xlsx
"""
import numpy as np
import openpyxl
import datetime

data = np.load(r'D:\数学建模\sol2_compliant.npz', allow_pickle=True)
records = data['records']
warmup = data['warmup_records']
cal_log = data['calibration_log']
pred_params = data['pred_params']

# 合并预热期和正式期（预热期1月不计入正式结果，但需要连续SOC）
all_records = list(warmup) + list(records)
formal = list(records)  # 正式区间334天

total_plan = float(data['total_plan'])
total_emg = float(data['total_emg'])
total_g = float(data['total_g'])
total_r = float(data['total_r'])
total_w = float(data['total_w'])
days_emg = int(data['days_emg'])

print(f"预测参数: 负载({pred_params[0]},{pred_params[1]}) 光伏({pred_params[2]},{pred_params[3]})")
print(f"合计费用: {total_plan+total_emg:.2f}")
print(f"正式记录数: {len(formal)}")

# 写入Excel
wb = openpyxl.load_workbook(r'D:\数学建模\result2.xlsx')

# 计划购电量
ws = wb['计划购电量']
for i, rec in enumerate(formal):
    row = i + 2
    ws.cell(row, 1).value = rec['date']
    g = rec['g']
    for t in range(144):
        ws.cell(row, t + 2).value = round(float(g[t]), 4)
    ws.cell(row, 146).value = round(float(g.sum()), 2)
    ws.cell(row, 147).value = round(float(rec['cost_plan']), 2)
print("计划购电量 sheet 完成")

# 充放电量
ws2 = wb['充放电量']
time_slots = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00', '16:00-20:00', '20:00-24:00']
for i, rec in enumerate(formal):
    base_row = i * 6 + 2
    c = rec['c']; d = rec['d']; E0 = rec['E0']; E_end = rec['E_end']
    for k in range(6):
        row = base_row + k
        ws2.cell(row, 2).value = time_slots[k]
        c_sum = float(c[k*24:(k+1)*24].sum())
        d_sum = float(d[k*24:(k+1)*24].sum())
        ws2.cell(row, 3).value = round(c_sum, 4)
        ws2.cell(row, 4).value = round(d_sum, 4)
        if k == 0:
            ws2.cell(row, 1).value = rec['date']
            ws2.cell(row, 5).value = datetime.time(0, 0)
            ws2.cell(row, 6).value = round(float(E0), 2)
        elif k == 1:
            ws2.cell(row, 5).value = '24:00'
            ws2.cell(row, 6).value = round(float(E_end), 2)
        else:
            ws2.cell(row, 5).value = None
            ws2.cell(row, 6).value = None
print("充放电量 sheet 完成")

# 紧急购电量
ws3 = wb['紧急购电量']
if ws3.max_row > 1:
    ws3.delete_rows(2, ws3.max_row - 1)
emg_events = []
for rec in formal:
    r = rec['r']
    date = rec['date']
    t = 0
    while t < 144:
        if r[t] > 1e-6:
            start_t = t
            emg_sum = 0.0
            while t < 144 and r[t] > 1e-6:
                emg_sum += r[t]
                t += 1
            end_t = t - 1
            start_min = start_t * 10
            end_min = (end_t + 1) * 10
            start_label = f"{start_min//60}:{start_min%60:02d}"
            end_label = "0:00+1" if end_min >= 1440 else f"{end_min//60}:{end_min%60:02d}"
            emg_events.append((date, f"{start_label}-{end_label}", round(float(emg_sum), 2)))
        else:
            t += 1
for i, (date, time_label, amount) in enumerate(emg_events):
    row = i + 2
    ws3.cell(row, 1).value = date
    ws3.cell(row, 2).value = time_label
    ws3.cell(row, 3).value = amount
print(f"紧急购电量 sheet 完成 ({len(emg_events)}个事件)")

wb.save(r'D:\数学建模\result2_compliant.xlsx')
print("\nresult2_compliant.xlsx 保存完成！")

# 打印对比
print("\n" + "="*60)
print("三版本对比")
print("="*60)
print(f"{'指标':<15} {'非合规版':>15} {'严格合规版':>15} {'论文版':>15}")
print("-"*60)
print(f"{'总费用(元)':<15} {13909376.66:>15.2f} {total_plan+total_emg:>15.2f} {14021565.12:>15.2f}")
print(f"{'计划购电(kWh)':<15} {21387992.67:>15.2f} {total_g:>15.2f} {21701732.99:>15.2f}")
print(f"{'紧急购电(kWh)':<15} {151486.53:>15.2f} {total_r:>15.2f} {124045.19:>15.2f}")
print(f"{'紧急费占比':<15} {'5.71%':>15} {total_emg/(total_plan+total_emg)*100:.2f}%{'':>10} {'4.73%':>15}")
print(f"{'弃电(kWh)':<15} {2399159.91:>15.2f} {total_w:>15.2f} {'-':>15}")
print(f"{'紧急天数':<15} {'129/334':>15} {f'{days_emg}/334':>15} {'160/334':>15}")
print("="*60)
print(f"\n合规成本: {(total_plan+total_emg) - 13909376.66:.2f} 元 (比非合规版贵)")
print(f"对比论文: {14021565.12 - (total_plan+total_emg):.2f} 元 (比论文省)")
