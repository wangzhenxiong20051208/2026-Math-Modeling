# -*- coding: utf-8 -*-
"""
2026 CUMCM C题 问题三 — 完整求解脚本（一步到位输出result3.xlsx）
口径1: 总费用 = p*gP + 1.5p*up - 0.5p*down + 5p*r
合规: 参数用1月预热期标定，固定后跑2-12月，预测用滚动历史误差
"""
import numpy as np
import pandas as pd
from scipy.optimize import linprog
import openpyxl
import datetime

# ==================== 常量 ====================
N = 144
DT = 1 / 6
ETA = 0.9
EMIN = 1200.0
EMAX = 10800.0
E_INIT = 6000.0
PMAX = 5000.0
M = PMAX * DT
N_DAY = 365
REPORT_START = 31
ADJ_TIMES = {'0:00': 0, '6:00': 36, '12:00': 72, '18:00': 108}

# ==================== 数据加载 ====================
print("=" * 60)
print("步骤1: 加载数据")
print("=" * 60)
df1 = pd.read_excel(r'01_题目\C题\附件\附件1.xlsx')
price = df1['电价'].astype(float).values

load_df = pd.read_excel(r'01_题目\C题\附件\附件2.xlsx', sheet_name='小区负载')
pv_df = pd.read_excel(r'01_题目\C题\附件\附件2.xlsx', sheet_name='光伏发电实际功率')
dates = pd.to_datetime(load_df.iloc[:, 0]).tolist()
load_e = load_df.iloc[:, 1:].astype(float).values * DT
pv_e = pv_df.iloc[:, 1:].astype(float).values * DT

fc_df = pd.read_excel(r'01_题目\C题\附件\附件3.xlsx')
fc_df['日期'] = fc_df['日期'].ffill()
fc_df['日期'] = pd.to_datetime(fc_df['日期'])
fc_cols = [f'预报{i}小时' for i in range(1, 25)]
forecast = {}
for day_idx in range(N_DAY):
    day_date = dates[day_idx]
    day_fc = fc_df[fc_df['日期'] == day_date]
    forecast[day_idx] = {}
    for _, row in day_fc.iterrows():
        adj_time = row['预报时刻']
        if adj_time not in ADJ_TIMES:
            continue
        start_t = ADJ_TIMES[adj_time]
        hourly_kw = row[fc_cols].astype(float).values
        pv_10min_e = np.zeros(N)
        for h in range(24):
            t_base = start_t + h * 6
            if t_base >= N:
                break
            for k in range(6):
                t = t_base + k
                if t < N:
                    pv_10min_e[t] = hourly_kw[h] * DT
        forecast[day_idx][adj_time] = pv_10min_e

# 预报误差表
fc_err = {}
for adj_time in ADJ_TIMES:
    fc_err[adj_time] = np.full((N_DAY, N), np.nan)
for day_idx in range(1, N_DAY):
    for adj_time in ADJ_TIMES:
        if adj_time in forecast[day_idx]:
            fc_err[adj_time][day_idx] = forecast[day_idx][adj_time] - pv_e[day_idx]

print(f"  电价: {price.shape}, 负载: {load_e.shape}, 光伏: {pv_e.shape}")
print(f"  预报: {len(forecast)}天, 误差表已构建")

# ==================== 核心函数 ====================
def safety_margin(day_idx, adj_time, alpha, window=28):
    """滚动历史预报误差的alpha分位数"""
    if day_idx == 0:
        return np.zeros(N)
    lo = max(0, day_idx - window)
    hist = fc_err[adj_time][lo:day_idx, :]
    valid = ~np.isnan(hist).all(axis=1)
    hist = hist[valid]
    k = hist.shape[0]
    if k >= 8:
        return np.nanquantile(hist, alpha, axis=0)
    if k >= 3:
        pooled = np.empty((k, N))
        for t in range(N):
            a, b = max(0, t - 2), min(N, t + 3)
            pooled[:, t] = np.nanmean(hist[:, a:b], axis=1)
        return np.nanquantile(pooled, alpha, axis=0)
    return np.zeros(N)


def solve_lp(net_load, e0, t_start=0, t_end=N):
    """0:00计划层LP: min 购电费, 含储能约束"""
    n_opt = t_end - t_start
    if n_opt <= 0:
        return None
    iG, iC, iD, iS, iE = 0, n_opt, 2 * n_opt, 3 * n_opt, 4 * n_opt
    nv = 5 * n_opt
    c_obj = np.zeros(nv)
    for k in range(n_opt):
        c_obj[iG + k] = price[t_start + k]
    A_eq = np.zeros((2 * n_opt, nv))
    b_eq = np.zeros(2 * n_opt)
    for k in range(n_opt):
        t = t_start + k
        A_eq[k, iG + k] = 1
        A_eq[k, iD + k] = 1
        A_eq[k, iC + k] = -1
        A_eq[k, iS + k] = -1
        b_eq[k] = net_load[t]
        A_eq[n_opt + k, iE + k] = 1
        if k > 0:
            A_eq[n_opt + k, iE + k - 1] = -1
        A_eq[n_opt + k, iC + k] = -ETA
        A_eq[n_opt + k, iD + k] = 1.0 / ETA
        if k == 0:
            b_eq[n_opt + k] = e0
    bounds = ([(0, None)] * n_opt + [(0, M)] * n_opt + [(0, M)] * n_opt
              + [(0, None)] * n_opt + [(EMIN, EMAX)] * n_opt)
    res = linprog(c_obj, A_eq=A_eq, b_eq=b_eq, bounds=bounds, method='highs')
    if not res.success:
        return None
    return res.x[iG:iG + n_opt].copy(), res.x[iE:iE + n_opt].copy()


def solve_adjustment(plan_g, net_load, sm_val, e0, t_start, up_limit, t_end=N):
    """调整层LP: min 购电费+调整成本(口径1), 允许下调和有限上调"""
    n_opt = t_end - t_start
    if n_opt <= 0:
        return None
    iG, iC, iD, iS, iE = 0, n_opt, 2 * n_opt, 3 * n_opt, 4 * n_opt
    iDown, iUp = 5 * n_opt, 6 * n_opt
    nv = 7 * n_opt
    # 口径1: 目标 = p*gA + 0.5p*down + 0.5p*up = p*gP - 0.5p*down + 1.5p*up
    c_obj = np.zeros(nv)
    for k in range(n_opt):
        t = t_start + k
        c_obj[iG + k] = price[t]
        c_obj[iDown + k] = 0.5 * price[t]
        c_obj[iUp + k] = 0.5 * price[t]
    A_eq = np.zeros((2 * n_opt, nv))
    b_eq = np.zeros(2 * n_opt)
    for k in range(n_opt):
        t = t_start + k
        A_eq[k, iG + k] = 1
        A_eq[k, iD + k] = 1
        A_eq[k, iC + k] = -1
        A_eq[k, iS + k] = -1
        b_eq[k] = net_load[t] + sm_val[t]
        A_eq[n_opt + k, iE + k] = 1
        if k > 0:
            A_eq[n_opt + k, iE + k - 1] = -1
        A_eq[n_opt + k, iC + k] = -ETA
        A_eq[n_opt + k, iD + k] = 1.0 / ETA
        if k == 0:
            b_eq[n_opt + k] = e0
    A_ub = np.zeros((2 * n_opt, nv))
    b_ub = np.zeros(2 * n_opt)
    for k in range(n_opt):
        t = t_start + k
        A_ub[2 * k, iDown + k] = -1
        A_ub[2 * k, iG + k] = -1
        b_ub[2 * k] = -plan_g[t]
        A_ub[2 * k + 1, iUp + k] = -1
        A_ub[2 * k + 1, iG + k] = 1
        b_ub[2 * k + 1] = plan_g[t]
    bounds_g = [(0, (1 + up_limit) * plan_g[t_start + k]) for k in range(n_opt)]
    bounds = (bounds_g + [(0, M)] * n_opt + [(0, M)] * n_opt + [(0, None)] * n_opt
              + [(EMIN, EMAX)] * n_opt + [(0, None)] * n_opt + [(0, None)] * n_opt)
    res = linprog(c_obj, A_ub=A_ub, b_ub=b_ub, A_eq=A_eq, b_eq=b_eq, bounds=bounds, method='highs')
    if not res.success:
        return None
    return res.x[iG:iG + n_opt].copy(), res.x[iE:iE + n_opt].copy()


def execute_day(final_g, E_ref, l_act, v_act, e0, rho):
    """实时执行层: 贪心充放电, 储备线保护"""
    c = np.zeros(N)
    d = np.zeros(N)
    r = np.zeros(N)
    w = np.zeros(N)
    E = np.zeros(N)
    soc = e0
    for t in range(N):
        b = final_g[t] + v_act[t] - l_act[t]
        R = EMIN + rho * (E_ref[t] - EMIN) if E_ref is not None else EMIN
        if b >= 0:
            cmax = min(b, M, (EMAX - soc) / ETA)
            c[t] = max(0, cmax)
            w[t] = b - c[t]
        else:
            head = max(0.0, soc - R)
            dmax = min(-b, M, ETA * head)
            d[t] = max(0, dmax)
            r[t] = max(0, -b - d[t])
        soc = soc + ETA * c[t] - d[t] / ETA
        E[t] = soc
    plan_cost = float(np.sum(price * final_g))
    emg_cost = float(np.sum(5.0 * price * r))
    return c, d, r, w, E, plan_cost, emg_cost


def run_days(start_day, end_day, alpha, rho, up_limit, e_init):
    """运行指定日期区间，返回记录列表和期末储能"""
    e = e_init
    records = []
    for day_idx in range(start_day, end_day):
        pv_fc_0 = forecast[day_idx].get('0:00', np.zeros(N))
        sm_0 = safety_margin(day_idx, '0:00', alpha)
        net_load_0 = load_e[day_idx] - pv_fc_0 + sm_0
        res0 = solve_lp(net_load_0, e, 0, N)
        if res0 is None:
            plan_g = np.maximum(net_load_0, 0)
            E_plan = np.full(N, e)
        else:
            plan_g, E_plan = res0
        current_g = plan_g.copy()
        current_E = E_plan.copy()
        for adj_t in ['6:00', '12:00', '18:00']:
            if adj_t in forecast[day_idx]:
                pv_fc = forecast[day_idx][adj_t]
                t_start = ADJ_TIMES[adj_t]
                sm_val = safety_margin(day_idx, adj_t, alpha)
                net_load = load_e[day_idx] - pv_fc
                e_cur = current_E[t_start - 1]
                res = solve_adjustment(current_g, net_load, sm_val, e_cur, t_start, up_limit, N)
                if res is not None:
                    adj_g, E_adj = res
                    current_g[t_start:] = adj_g
                    current_E[t_start:] = E_adj
        c, d, r, w, E_act, plan_cost, emg_cost = execute_day(
            current_g, current_E, load_e[day_idx], pv_e[day_idx], e, rho)
        # 口径1调整成本: 0.5p*down + 0.5p*up (因为plan_cost=p*gA)
        adj_cost = 0.0
        for t in range(N):
            if current_g[t] < plan_g[t] - 1e-9:
                adj_cost += 0.5 * price[t] * (plan_g[t] - current_g[t])
            elif current_g[t] > plan_g[t] + 1e-9:
                adj_cost += 0.5 * price[t] * (current_g[t] - plan_g[t])
        rec = {
            'n': day_idx, 'date': dates[day_idx],
            'plan_g': plan_g, 'final_g': current_g,
            'c': c, 'd': d, 'r': r, 'w': w, 'E': E_act,
            'E0': e, 'E_end': E_act[-1],
            'plan_cost': plan_cost, 'adj_cost': adj_cost, 'emg_cost': emg_cost,
            'cost_total': plan_cost + adj_cost + emg_cost,
        }
        records.append(rec)
        e = E_act[-1]
    return records, e


# ==================== 步骤2: 1月参数标定 ====================
print("\n" + "=" * 60)
print("步骤2: 1月预热期参数标定（口径1）")
print("=" * 60)

# 第一步: 固定up=0.15, 搜索alpha x rho
print("  2a. 搜索 alpha x rho (up=0.15)...")
best_cost = float('inf')
best_ar = None
for alpha in [0.75, 0.80, 0.85, 0.90]:
    for rho in [0.55, 0.60, 0.65, 0.70]:
        recs, _ = run_days(0, 31, alpha, rho, 0.15, E_INIT)
        cost = sum(r['cost_total'] for r in recs)
        print(f"    alpha={alpha:.2f} rho={rho:.2f}: {cost:.2f}")
        if cost < best_cost:
            best_cost = cost
            best_ar = (alpha, rho)

print(f"  最优 alpha={best_ar[0]:.2f} rho={best_ar[1]:.2f}, 1月费用={best_cost:.2f}")

# 第二步: 固定最优alpha rho, 搜索up_limit
print(f"\n  2b. 搜索 up_limit (alpha={best_ar[0]:.2f}, rho={best_ar[1]:.2f})...")
best_up = None
for up in [0.05, 0.10, 0.15, 0.20, 0.30]:
    recs, _ = run_days(0, 31, best_ar[0], best_ar[1], up, E_INIT)
    cost = sum(r['cost_total'] for r in recs)
    print(f"    up={up:.2f}: {cost:.2f}")
    if cost < best_cost:
        best_cost = cost
        best_up = up

ALPHA_OPT, RHO_OPT, UP_OPT = best_ar[0], best_ar[1], best_up
print(f"\n  最终1月最优参数: alpha={ALPHA_OPT:.2f} rho={RHO_OPT:.2f} up={UP_OPT:.2f}")
print(f"  1月最优总费用: {best_cost:.2f}")

# ==================== 步骤3: 正式区间求解 ====================
print("\n" + "=" * 60)
print(f"步骤3: 正式区间求解 (alpha={ALPHA_OPT}, rho={RHO_OPT}, up={UP_OPT})")
print("=" * 60)

# 先跑1月得到期末储能
_, e_jan_end = run_days(0, 31, ALPHA_OPT, RHO_OPT, UP_OPT, E_INIT)
print(f"  1月期末储电量: {e_jan_end:.2f} kWh")

# 跑正式区间
all_records, e_final = run_days(31, N_DAY, ALPHA_OPT, RHO_OPT, UP_OPT, e_jan_end)
print(f"  正式区间天数: {len(all_records)}")
print(f"  期末储电量: {e_final:.2f} kWh")

# ==================== 步骤4: 汇总与验证 ====================
print("\n" + "=" * 60)
print("步骤4: 汇总与口径1验证")
print("=" * 60)

formal = all_records
total_plan_g = sum(float(r['plan_g'].sum()) for r in formal)
total_g = sum(float(r['final_g'].sum()) for r in formal)
total_r = sum(float(r['r'].sum()) for r in formal)
total_w = sum(float(r['w'].sum()) for r in formal)
total_plan_cost = sum(r['plan_cost'] for r in formal)
total_adj_cost = sum(r['adj_cost'] for r in formal)
total_emg_cost = sum(r['emg_cost'] for r in formal)
total_cost = total_plan_cost + total_adj_cost + total_emg_cost
days_emg = sum(1 for r in formal if r['r'].sum() > 1e-6)
days_adj = sum(1 for r in formal if r['adj_cost'] > 1e-6)

# 口径1独立验证
total_p_gP = sum(float(np.sum(price * r['plan_g'])) for r in formal)
total_1p5up = 0.0
total_0p5down = 0.0
for r in formal:
    diff = r['final_g'] - r['plan_g']
    total_1p5up += float(np.sum(1.5 * price * np.maximum(diff, 0)))
    total_0p5down += float(np.sum(0.5 * price * np.maximum(-diff, 0)))
koujing1_total = total_p_gP + total_1p5up - total_0p5down + total_emg_cost

print(f"  0:00计划购电量: {total_plan_g:.2f} kWh")
print(f"  最终执行购电量: {total_g:.2f} kWh")
print(f"  紧急购电量: {total_r:.2f} kWh")
print(f"  弃用电量: {total_w:.2f} kWh")
print(f"  实际购电费 p*gA: {total_plan_cost:.2f}")
print(f"  调整成本: {total_adj_cost:.2f}")
print(f"  紧急购电费: {total_emg_cost:.2f}")
print(f"  合计费用: {total_cost:.2f}")
print(f"  紧急费占比: {total_emg_cost/total_cost*100:.2f}%")
print(f"  发生紧急购电天数: {days_emg}/334")
print(f"  发生调整天数: {days_adj}/334")
print(f"\n  口径1独立验证:")
print(f"    p*gP = {total_p_gP:.2f}")
print(f"    1.5p*up = {total_1p5up:.2f}")
print(f"    0.5p*down = {total_0p5down:.2f}")
print(f"    5p*r = {total_emg_cost:.2f}")
print(f"    口径1总费 = {koujing1_total:.2f}")
print(f"    实际计算总费 = {total_cost:.2f}")
print(f"    差异 = {koujing1_total - total_cost:.2f} (应接近0)")

# ==================== 步骤5: 写入result3.xlsx ====================
print("\n" + "=" * 60)
print("步骤5: 写入result3.xlsx")
print("=" * 60)

wb = openpyxl.load_workbook(r'01_题目\C题\附件\附件5\result3.xlsx')

# Sheet1: 计划购电量
ws = wb['计划购电量']
if ws.max_row > 1:
    ws.delete_rows(2, ws.max_row - 1)
for i, rec in enumerate(formal):
    row = i + 2
    ws.cell(row, 1).value = rec['date']
    g = rec['plan_g']
    for t in range(144):
        ws.cell(row, t + 2).value = round(float(g[t]), 4)
    ws.cell(row, 146).value = round(float(g.sum()), 2)
    ws.cell(row, 147).value = round(float(np.sum(price * g)), 2)
print("  计划购电量 sheet 完成")

# Sheet2: 调整购电量
ws2 = wb['调整购电量']
if ws2.max_row > 1:
    ws2.delete_rows(2, ws2.max_row - 1)
for i, rec in enumerate(formal):
    row = i + 2
    ws2.cell(row, 1).value = rec['date']
    g = rec['final_g']
    plan_g = rec['plan_g']
    for t in range(144):
        ws2.cell(row, t + 2).value = round(float(g[t]), 4)
    ws2.cell(row, 146).value = round(float(g.sum()), 2)
    # 全天购电费 = 合同费(口径1) + 紧急费
    diff = g - plan_g
    up = np.maximum(diff, 0)
    down = np.maximum(-diff, 0)
    contract_cost = float(np.sum(price * plan_g) + np.sum(1.5 * price * up) - np.sum(0.5 * price * down))
    emg_cost = float(np.sum(5.0 * price * rec['r']))
    ws2.cell(row, 147).value = round(contract_cost + emg_cost, 2)
print("  调整购电量 sheet 完成")

# Sheet3: 充放电量
ws3 = wb['充放电量']
if ws3.max_row > 1:
    ws3.delete_rows(2, ws3.max_row - 1)
time_slots = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00', '16:00-20:00', '20:00-24:00']
for i, rec in enumerate(formal):
    base_row = i * 6 + 2
    c = rec['c']
    d = rec['d']
    E0 = rec['E0']
    E_end = rec['E_end']
    for k in range(6):
        row = base_row + k
        ws3.cell(row, 2).value = time_slots[k]
        c_sum = float(c[k * 24:(k + 1) * 24].sum())
        d_sum = float(d[k * 24:(k + 1) * 24].sum())
        ws3.cell(row, 3).value = round(c_sum, 4)
        ws3.cell(row, 4).value = round(d_sum, 4)
        if k == 0:
            ws3.cell(row, 1).value = rec['date']
            ws3.cell(row, 5).value = datetime.time(0, 0)
            ws3.cell(row, 6).value = round(float(E0), 2)
        elif k == 1:
            ws3.cell(row, 5).value = '24:00'
            ws3.cell(row, 6).value = round(float(E_end), 2)
print("  充放电量 sheet 完成")

# Sheet4: 紧急购电量
ws4 = wb['紧急购电量']
if ws4.max_row > 1:
    ws4.delete_rows(2, ws4.max_row - 1)
emg_events = []
for rec in formal:
    r = rec['r']
    date = rec['date']
    t = 0
    while t < 144:
        if r[t] > 1e-6:
            start_t = t
            emg_sum = 0.0
            while t < 144 and r[t] > 1e-6:
                emg_sum += r[t]
                t += 1
            end_t = t - 1
            start_min = start_t * 10
            end_min = (end_t + 1) * 10
            start_label = f"{start_min // 60}:{start_min % 60:02d}"
            end_label = "0:00+1" if end_min >= 1440 else f"{end_min // 60}:{end_min % 60:02d}"
            emg_events.append((date, f"{start_label}-{end_label}", round(float(emg_sum), 2)))
        else:
            t += 1
for i, (date, time_label, amount) in enumerate(emg_events):
    row = i + 2
    ws4.cell(row, 1).value = date
    ws4.cell(row, 2).value = time_label
    ws4.cell(row, 3).value = amount
print(f"  紧急购电量 sheet 完成 ({len(emg_events)}个事件)")

wb.save(r'result3.xlsx')

# ==================== 完成 ====================
print("\n" + "=" * 60)
print("完成！")
print("=" * 60)
print(f"  输出文件: result3.xlsx")
print(f"  最终总费用(口径1): {total_cost:.2f} 元")
print(f"  标定参数: alpha={ALPHA_OPT}, rho={RHO_OPT}, up_limit={UP_OPT}")
print(f"  合规性: 1月标定 -> 固定参数 -> 2-12月求解，无事后选参")
